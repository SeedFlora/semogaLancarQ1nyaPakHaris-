"""End-to-end test of the three-stage pipeline through its command-line interfaces.

synthetic artifacts -> export_per_sample_predictions.py -> analyze_uncertainty.py
-> fill_manuscript.py on a mock manuscript holding every CONTRACT.md section 4 token.

The unit tests of each stage use their own synthetic inputs; this is the only test in which the
real output of one stage is the input of the next, so it guards the file contracts between them
(CSV schema and manifest -> analysis; fill_values.json and PNG sizes -> fill).

Run with:  python -m pytest tests/test_e2e_chain.py -q
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
for p in (ROOT, ROOT / "tests"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

import e2e_docx

import analyze_uncertainty as au
import fill_manuscript as fm
import make_synthetic_fixture as fx

TOKEN_RE = re.compile(r"⟦([A-Z0-9_]+)(?::[^⟧]*)?⟧")
ASCII_MINUS_RE = re.compile(r"(?<![\w.])-(?=\d)")
DEPLOYED = "dinov3_large__mE5_large"
FIG_PX = {"fig_reliability.png": (2100, 780), "fig_selective.png": (2100, 840)}
WP_DOCPR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"


def _run(script: str, *args: object) -> subprocess.CompletedProcess:
    """Run one stage exactly as a user would (separate interpreter, piped output)."""
    return subprocess.run([sys.executable, str(ROOT / script), *map(str, args)], check=False,
                          capture_output=True, text=True, encoding="utf-8", errors="replace",
                          timeout=600)


@pytest.fixture(scope="module")
def chain(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Path]:
    base = tmp_path_factory.mktemp("e2e")
    artifacts = base / "artifacts"
    fx.build_fixture(artifacts, n_samples=900, dim=16, verbose=False)   # all three default pairs
    preds, results = base / "per_sample_predictions", base / "results"

    exp = _run("export_per_sample_predictions.py", "--artifacts", artifacts, "--out", preds, "--no-expect")
    assert exp.returncode == 0, exp.stdout + exp.stderr
    ana = _run("analyze_uncertainty.py", "--preds", preds, "--out", results)
    assert ana.returncode == 0, ana.stdout + ana.stderr

    docx_in, docx_out = base / "all_keys.docx", base / "filled.docx"
    e2e_docx.build_all_keys_docx(docx_in, base / "assets")
    fill = _run("fill_manuscript.py", "--docx", docx_in, "--results", results, "--out", docx_out,
                "--report-json", base / "fill_report.json")
    assert fill.returncode == 0, fill.stdout + fill.stderr
    return {"preds": preds, "results": results, "docx_in": docx_in, "docx_out": docx_out,
            "report": base / "fill_report.json"}


def _values(results: Path) -> dict[str, str]:
    return json.loads((results / "fill_values.json").read_text(encoding="utf-8"))


def test_key_lists_agree_with_contract():
    keys = e2e_docx.contract_keys()
    assert len(keys) == 84 and len(set(keys)) == 84
    assert au.REQUIRED_KEYS == keys
    assert fm.contract_keys() == keys


def test_stage1_to_stage2_contract(chain):
    preds, results = chain["preds"], chain["results"]
    manifest = json.loads((preds / "manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["files"]) == 12
    for entry in manifest["files"]:
        assert (preds / entry["file"]).is_file()
    for name in ("results.json", "fill_values.json", "tables.md", "claims_check.md",
                 *FIG_PX):
        assert (results / name).is_file(), name
    for name, px in FIG_PX.items():
        with Image.open(results / name) as im:
            assert im.size == px, name

    # numbers carried across the stage boundary agree
    values = _values(results)
    by = {(e["image"] + "__" + e["text"], e["checkpoint"], e["mode"]): e for e in manifest["files"]}
    for row, key in (("R1", ("cb", "argmax")), ("R2", ("cb", "pgs")), ("R3", ("pgs", "argmax")),
                     ("R4", ("pgs", "pgs"))):
        acc = by[(DEPLOYED, *key)]["metrics"]["test"]["accuracy"]
        assert values[f"T5_{row}_ACC"] == f"{acc:.4f}", row
    b, c = (int(x) for x in values["T6_R4_BC"].split(" / "))
    assert b + c <= manifest["top1_disagreements_pgs_vs_argmax"]
    res = json.loads((results / "results.json").read_text(encoding="utf-8"))
    assert res["reproduction"]["top1_disagreements_pgs_vs_argmax"] == manifest["top1_disagreements_pgs_vs_argmax"]

    # --no-expect export: C5 cannot be checked, and the claims file says what to do about it
    claims = (results / "claims_check.md").read_text(encoding="utf-8")
    assert "*Action:* None" not in claims
    assert "without --no-expect" in claims


def test_fill_values_formatting(chain):
    values = _values(chain["results"])
    assert set(values) == set(e2e_docx.contract_keys())
    for key, text in values.items():
        assert isinstance(text, str) and text.strip(), key
        assert not ASCII_MINUS_RE.search(text), (key, text)
        assert "⟦" not in text and "⟧" not in text, key
    assert re.fullmatch(r"\d+ / \d+", values["T6_R1_BC"])
    assert re.fullmatch(r"0\.\d{4} \[0\.\d{4}, 0\.\d{4}\]", values["T5_R1_ECE"])
    assert re.fullmatch(r"[+−]\d+\.\d{2} \[[+−]?\d+\.\d{2}, [+−]?\d+\.\d{2}\]", values["T6_R5_DF1"])


def test_filled_docx_has_no_tokens_and_every_value(chain):
    values = _values(chain["results"])
    out = chain["docx_out"]
    with zipfile.ZipFile(chain["docx_in"]) as zin:
        before = TOKEN_RE.findall(zin.read("word/document.xml").decode("utf-8"))
    assert sorted(before) == sorted(e2e_docx.contract_keys())

    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                text = z.read(name).decode("utf-8", "replace")
                assert not TOKEN_RE.search(text), name
                assert "⟦" not in text and "⟧" not in text, name

    doc = Document(str(out))                     # opens with python-docx
    t1, t5, t6, t7 = doc.tables
    for i, key in enumerate(e2e_docx.T1_KEYS):
        assert t1.cell(i, 1).text == values[key], key
    for table, prefix, cols, n_rows in ((t5, "T5", e2e_docx.T5_COLS, 5), (t6, "T6", e2e_docx.T6_COLS, 5),
                                        (t7, "T7", e2e_docx.T7_COLS, 6)):
        for r in range(1, n_rows + 1):
            for j, col in enumerate(cols, start=1):
                key = f"{prefix}_R{r}_{col}"
                assert table.cell(r, j).text == values[key], key
    paragraphs = [p.text for p in doc.paragraphs]
    assert f"Fitted temperature T = {values['T5_TEMP']}." in paragraphs
    for key in e2e_docx.TXT_KEYS:
        assert e2e_docx.TXT_LEAD.get(key, "") + values[key] in paragraphs, key

    runs = [r for p in doc.paragraphs for r in p.runs]
    runs += [r for t in doc.tables for row in t.rows for cell in row.cells for p in cell.paragraphs
             for r in p.runs]
    assert not [r.text for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]

    report = json.loads(chain["report"].read_text(encoding="utf-8"))
    assert report["ok"] and report["n_tokens_replaced"] == 84
    assert not report["missing"] and not report["unknown_keys"] and not report["contract_keys_absent"]
    assert not report["errors"] and not report["warnings"] and not report["missing_figures"]


def test_figures_replaced_with_stage2_pngs(chain):
    results, out = chain["results"], chain["docx_out"]
    doc = Document(str(out))
    names = [el.get("name") for el in doc.element.body.iter(WP_DOCPR)]
    assert names == ["Figure 5", "Figure 6"]
    media = {}
    for shape in doc.inline_shapes:
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        media[shape._inline.docPr.get("name")] = doc.part.related_parts[rid].blob
    for name, png in (("Figure 5", "fig_reliability.png"), ("Figure 6", "fig_selective.png")):
        assert hashlib.sha256(media[name]).digest() == hashlib.sha256((results / png).read_bytes()).digest()
