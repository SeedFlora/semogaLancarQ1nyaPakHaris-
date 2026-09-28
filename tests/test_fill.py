"""Tests for fill_manuscript.py (stage 3, CONTRACT.md §4).

The main fixture builds a small DOCX with python-docx (plus lxml injections for
the tracked insertion, a split token and the placeholder docPr names), fills it
and keeps both files in ``_work/fill_test/`` for manual inspection.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import fill_manuscript as fm

WORK = ROOT / "_work" / "fill_test"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {
    "w": W,
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
FIG5, FIG6 = "PLACEHOLDER_FIG5_RELIABILITY", "PLACEHOLDER_FIG6_SELECTIVE"

SPECIFIC_VALUES = {
    "T5_R3_ACC": "0.8116",
    "TXT_CALIBRATION": ("At the fixed PGS checkpoint, averaging lowers ECE from 0.0150 to 0.0123 "
                        "(overlapping CIs) & Brier < 0.30. Temperature scaling (T = 1.23) "
                        "reduces ECE further."),
    "T6_R1_BC": "84 / 45",
    "T6_R1_P": "< 0.001",
    "T6_R1_DACC": "−0.43 [−0.80, −0.05]",
    "TXT_REPRO": "to the fourth decimal place for all four checkpoint–mode combinations.",
    "T7_R5_AUROC": "0.7012 [0.6900, 0.7100]",
    "T7_R1_AUROC": "0.8123 [0.8012, 0.8230]",
    "T6_R2_P": "0.0421",
    "T6_R3_P": "0.237",
    "T5_TEMP": "1.23",
    "T5_R1_ECE": "0.0123 [0.0101, 0.0150]",
    "T1_AUROC": "0.8123 vs. 0.7512",
}

# token -> expected count in the main test document
EXPECTED_COUNTS = {
    "T5_R3_ACC": 1, "TXT_CALIBRATION": 1, "T6_R1_BC": 1, "T6_R1_P": 1, "T6_R1_DACC": 1,
    "TXT_REPRO": 1, "T7_R5_AUROC": 1, "T7_R1_AUROC": 1, "T6_R2_P": 1, "T6_R3_P": 1,
    "T5_TEMP": 1, "T5_R1_ECE": 1, "T1_AUROC": 1,
}


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_png(path: Path, size: tuple[int, int], color: tuple[int, int, int]) -> bytes:
    Image.new("RGB", size, color).save(path, format="PNG")
    return path.read_bytes()


def fill_values() -> dict[str, str]:
    values = {k: f"v_{k.lower()}" for k in fm.contract_keys()}
    values.update(SPECIFIC_VALUES)
    values["ZZ_NOT_IN_DOC"] = "unused value"
    return values


def highlighted_run(paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def rename_docpr(inline_shape, name: str) -> None:
    inline_shape._inline.docPr.set("name", name)


def para_text(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(f"{{{W}}}t"))


def run_has_highlight(r_el) -> bool:
    return r_el.find("w:rPr/w:highlight", NS) is not None


def read_zip(path: Path) -> tuple[list[zipfile.ZipInfo], dict[str, bytes]]:
    with zipfile.ZipFile(path) as zf:
        infos = zf.infolist()
        return infos, {i.filename: zf.read(i.filename) for i in infos}


def rels_targets(blobs: dict[str, bytes], source: str) -> dict[str, str]:
    rels = etree.fromstring(blobs[fm.rels_part_name(source)])
    return {r.get("Id"): fm.resolve_target(source, r.get("Target"))
            for r in rels.iter(f"{{{NS['rel']}}}Relationship")
            if r.get("TargetMode") != "External"}


def figure_info(blobs: dict[str, bytes], docpr_name: str) -> list[dict]:
    """For each docPr with that name: media part, wp:extent and a:ext sizes."""
    doc = etree.fromstring(blobs["word/document.xml"])
    targets = rels_targets(blobs, "word/document.xml")
    found = []
    for docpr in doc.iter(f"{{{NS['wp']}}}docPr"):
        if docpr.get("name") != docpr_name:
            continue
        inline = docpr.getparent()
        blip = inline.find(".//a:blip", NS)
        extent = inline.find("wp:extent", NS)
        a_ext = inline.find(".//pic:spPr/a:xfrm/a:ext", NS)
        found.append({
            "media": targets[blip.get(f"{{{NS['r']}}}embed")],
            "cx": int(extent.get("cx")), "cy": int(extent.get("cy")),
            "a_cx": int(a_ext.get("cx")), "a_cy": int(a_ext.get("cy")),
        })
    return found


def write_results(results: Path, values: dict[str, str] | None = None,
                  figs: bool = True) -> dict[str, bytes]:
    results.mkdir(parents=True, exist_ok=True)
    (results / "fill_values.json").write_text(
        json.dumps(values if values is not None else fill_values(), ensure_ascii=False, indent=1),
        encoding="utf-8")
    pngs: dict[str, bytes] = {}
    if figs:
        pngs[FIG5] = make_png(results / "fig_reliability.png", (2100, 780), (30, 90, 200))
        pngs[FIG6] = make_png(results / "fig_selective.png", (2100, 840), (220, 120, 20))
    return pngs


def build_test_docx(path: Path, assets: Path) -> None:
    """The main test manuscript: every token situation named in the task."""
    assets.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Running head – AUROC ⟦T1_AUROC⟧"

    # P0: plain inline token in a normal (non-highlighted) run
    doc.add_paragraph("Accuracy of the served path was ⟦T5_R3_ACC⟧ on the test set.")
    # P1: highlighted run holding a token with a description
    p = doc.add_paragraph()
    p.add_run("Calibration. ")
    highlighted_run(p, "⟦TXT_CALIBRATION: two to four sentences on ECE, Brier and NLL⟧")
    # P2: several tokens in one w:t
    p = doc.add_paragraph()
    highlighted_run(p, "b / c = ⟦T6_R1_BC⟧ (p = ⟦T6_R1_P⟧), "
                       "Δacc = ⟦T6_R1_DACC⟧")
    # P3: highlighted run WITHOUT token (must stay highlighted)
    p = doc.add_paragraph()
    highlighted_run(p, "This highlighted note has no token.")
    # P4: token inside a tracked insertion <w:ins>
    p = doc.add_paragraph("The exported records reproduce the accuracy and macro-F1 values "
                          "of Table 4 ")
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "9001")
    ins.set(qn("w:author"), "Reviser")
    ins.set(qn("w:date"), "2026-09-28T00:00:00Z")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    hl = OxmlElement("w:highlight")
    hl.set(qn("w:val"), "yellow")
    rpr.append(hl)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "⟦TXT_REPRO: sentence fragment⟧"
    r.append(t)
    ins.append(r)
    p._p.append(ins)
    # P5: token split across two runs with a proofErr between them
    p = doc.add_paragraph()
    r1 = highlighted_run(p, "AUROC of MI: ⟦T7_R5_")
    highlighted_run(p, "AUROC⟧ overall.")
    proof = OxmlElement("w:proofErr")
    proof.set(qn("w:type"), "spellStart")
    r1._r.addnext(proof)
    # P6: token with description split across three runs
    p = doc.add_paragraph()
    p.add_run("Best: ")
    highlighted_run(p, "⟦T7_R1")
    highlighted_run(p, "_AUROC: area under")
    highlighted_run(p, " the ROC curve⟧")
    p.add_run(".")
    # P7: two split tokens sharing the middle run
    p = doc.add_paragraph()
    highlighted_run(p, "p-values: ⟦T6_R2_")
    highlighted_run(p, "P⟧ and ⟦T6_R3_")
    highlighted_run(p, "P⟧.")
    # P8: token whose w:t has a trailing space but lost xml:space="preserve"
    p = doc.add_paragraph()
    run = highlighted_run(p, "⟦T5_TEMP⟧ ")
    run._r.find(qn("w:t")).attrib.pop(XML_SPACE, None)
    p.add_run("is the fitted temperature.")
    # Table cell token
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Row"
    table.cell(0, 1).text = "ECE [95% CI]"
    table.cell(1, 0).text = "cb/argmax"
    highlighted_run(table.cell(1, 1).paragraphs[0], "⟦T5_R1_ECE⟧")
    # Figures: fig5 placeholder has ~0.5 % aspect difference (kept), fig6 a large one
    doc.add_paragraph("Figure 5 placeholder:")
    make_png(assets / "ph5.png", (2100, 784), (200, 200, 200))
    rename_docpr(doc.add_picture(str(assets / "ph5.png"), width=Inches(6.5)), FIG5)
    doc.add_paragraph("Figure 6 placeholder:")
    make_png(assets / "ph6.png", (2100, 600), (150, 150, 150))
    rename_docpr(doc.add_picture(str(assets / "ph6.png"), width=Inches(6.5)), FIG6)
    doc.save(str(path))


@pytest.fixture(scope="module")
def built() -> tuple[Path, Path, dict[str, bytes]]:
    WORK.mkdir(parents=True, exist_ok=True)
    src = WORK / "test_input.docx"
    build_test_docx(src, WORK / "assets")
    pngs = write_results(WORK / "results")
    return src, WORK / "results", pngs


def run_fill(*args: str) -> int:
    return fm.main([str(a) for a in args])


# ---------------------------------------------------------------------------
# main end-to-end test
# ---------------------------------------------------------------------------


def test_fill_end_to_end(built):
    src, results, pngs = built
    out = WORK / "test_filled.docx"
    report_json = WORK / "test_fill_report.json"
    before = sha256(src)
    assert run_fill("--docx", src, "--results", results, "--out", out,
                    "--report-json", report_json) == 0
    assert sha256(src) == before, "input must never be modified"

    in_infos, in_blobs = read_zip(src)
    out_infos, out_blobs = read_zip(out)
    values = fill_values()

    # -- all tokens replaced, in every XML part --------------------------------
    for name, data in out_blobs.items():
        if name.endswith(".xml"):
            text = data.decode("utf-8")
            assert "⟦" not in text and "⟧" not in text, name

    body = etree.fromstring(out_blobs["word/document.xml"]).find("w:body", NS)
    paras = body.findall("w:p", NS)
    texts = [para_text(p) for p in paras]
    assert texts[0] == f"Accuracy of the served path was {values['T5_R3_ACC']} on the test set."
    assert texts[1] == "Calibration. " + values["TXT_CALIBRATION"]
    assert texts[2] == (f"b / c = {values['T6_R1_BC']} (p = {values['T6_R1_P']}), "
                        f"Δacc = {values['T6_R1_DACC']}")
    assert texts[3] == "This highlighted note has no token."
    assert texts[4] == ("The exported records reproduce the accuracy and macro-F1 values of "
                        "Table 4 " + values["TXT_REPRO"])
    assert texts[5] == f"AUROC of MI: {values['T7_R5_AUROC']} overall."
    assert texts[6] == f"Best: {values['T7_R1_AUROC']}."
    assert texts[7] == f"p-values: {values['T6_R2_P']} and {values['T6_R3_P']}."
    assert texts[8] == f"{values['T5_TEMP']} is the fitted temperature."

    # -- highlight removed from filled runs, kept elsewhere ---------------------
    for i in (1, 2, 5, 6, 7, 8):
        assert not any(run_has_highlight(r) for r in paras[i].iter(f"{{{W}}}r")), i
    assert all(run_has_highlight(r) for r in paras[3].findall("w:r", NS))
    # xml:space preserve added where the result has a trailing space
    t8 = paras[8].find("w:r/w:t", NS)
    assert t8.text == values["T5_TEMP"] + " " and t8.get(XML_SPACE) == "preserve"
    # the value of a split token lands in the first run, the proofErr survives
    assert paras[5].find("w:proofErr", NS) is not None
    assert paras[5].findall("w:r", NS)[0].find("w:t", NS).text.endswith(values["T7_R5_AUROC"])

    # -- tracked insertion wrapper intact --------------------------------------
    ins_list = paras[4].findall("w:ins", NS)
    assert len(ins_list) == 1
    ins = ins_list[0]
    assert ins.get(f"{{{W}}}author") == "Reviser" and ins.get(f"{{{W}}}id") == "9001"
    ins_run = ins.find("w:r", NS)
    assert ins_run.find("w:t", NS).text == values["TXT_REPRO"]
    assert not run_has_highlight(ins_run)

    # -- table cell and header -------------------------------------------------
    tbl = body.find("w:tbl", NS)
    cell_texts = [para_text(p) for p in tbl.iter(f"{{{W}}}p")]
    assert values["T5_R1_ECE"] in cell_texts
    header_names = [n for n in out_blobs if n.startswith("word/header")]
    header_text = "".join(para_text(p) for n in header_names
                          for p in etree.fromstring(out_blobs[n]).iter(f"{{{W}}}p"))
    assert header_text == f"Running head – AUROC {values['T1_AUROC']}"

    # -- figures ---------------------------------------------------------------
    assert b"PLACEHOLDER_FIG" not in out_blobs["word/document.xml"]
    in5, in6 = figure_info(in_blobs, FIG5)[0], figure_info(in_blobs, FIG6)[0]
    out5, out6 = figure_info(out_blobs, "Figure 5"), figure_info(out_blobs, "Figure 6")
    assert len(out5) == 1 and len(out6) == 1
    out5, out6 = out5[0], out6[0]
    assert out_blobs[out5["media"]] == pngs[FIG5]
    assert out_blobs[out6["media"]] == pngs[FIG6]
    assert out5["media"] == in5["media"] and out6["media"] == in6["media"]  # in-place
    # fig5: 2100x784 -> 2100x780 is within the 1 % tolerance: extent untouched
    assert (out5["cx"], out5["cy"], out5["a_cy"]) == (in5["cx"], in5["cy"], in5["a_cy"])
    # fig6: 2100x600 -> 2100x840: width kept, height recomputed in both places
    assert out6["cx"] == in6["cx"] and out6["a_cx"] == in6["a_cx"]
    assert out6["cy"] == round(in6["cx"] * 840 / 2100) != in6["cy"]
    assert out6["a_cy"] == round(in6["a_cx"] * 840 / 2100)
    ct = etree.fromstring(out_blobs["[Content_Types].xml"])
    assert any(d.get("Extension").lower() == "png" for d in ct.findall("ct:Default", NS))

    # -- every other zip entry byte-identical, order and compression kept ------
    assert [i.filename for i in out_infos] == [i.filename for i in in_infos]
    assert [i.compress_type for i in out_infos] == [i.compress_type for i in in_infos]
    changed = {n for n in in_blobs if in_blobs[n] != out_blobs[n]}
    assert changed == {"word/document.xml", *header_names, in5["media"], in6["media"]}

    # -- report ----------------------------------------------------------------
    report = json.loads(report_json.read_text(encoding="utf-8"))
    assert report["ok"] is True
    assert report["replaced"] == EXPECTED_COUNTS
    assert report["missing"] == {}
    assert "ZZ_NOT_IN_DOC" in report["unknown_keys"]
    assert report["contract_keys_absent"] == []
    assert sorted(s["key"] for s in report["split_tokens"]) == [
        "T6_R2_P", "T6_R3_P", "T7_R1_AUROC", "T7_R5_AUROC"]
    assert all(s["replaced"] for s in report["split_tokens"])
    assert {f["final_name"] for f in report["figures"]} == {"Figure 5", "Figure 6"}
    assert report["errors"] == []

    # -- the result opens with python-docx -------------------------------------
    doc = Document(str(out))
    assert len(doc.inline_shapes) == 2
    assert doc.paragraphs[0].text.startswith("Accuracy of the served path was 0.8116")


# ---------------------------------------------------------------------------
# behaviour on incomplete input / options
# ---------------------------------------------------------------------------


def test_missing_key_exit_code_and_allow_missing(built, tmp_path, capsys):
    src, _, _ = built
    values = fill_values()
    del values["T5_R3_ACC"]            # single-run token
    del values["T7_R5_AUROC"]          # split token
    results = tmp_path / "results"
    write_results(results, values)
    out = tmp_path / "out.docx"
    report_json = tmp_path / "r.json"
    assert run_fill("--docx", src, "--results", results, "--out", out,
                    "--report-json", report_json) == 1
    assert "MISSING keys" in capsys.readouterr().out
    report = json.loads(report_json.read_text(encoding="utf-8"))
    assert report["missing"] == {"T5_R3_ACC": 1, "T7_R5_AUROC": 1}
    assert report["written"] is True
    _, blobs = read_zip(out)
    body = etree.fromstring(blobs["word/document.xml"]).find("w:body", NS)
    paras = body.findall("w:p", NS)
    assert "⟦T5_R3_ACC⟧" in para_text(paras[0])
    assert "⟦T7_R5_AUROC⟧" in para_text(paras[5])
    assert all(run_has_highlight(r) for r in paras[5].findall("w:r", NS))  # left for the author
    assert "⟦" not in para_text(paras[1])                            # others still filled

    assert run_fill("--docx", src, "--results", results, "--out", out, "--allow-missing") == 0


def test_keep_highlight(built, tmp_path):
    src, results, _ = built
    out = tmp_path / "kept.docx"
    assert run_fill("--docx", src, "--results", results, "--out", out, "--keep-highlight") == 0
    _, blobs = read_zip(out)
    body = etree.fromstring(blobs["word/document.xml"]).find("w:body", NS)
    paras = body.findall("w:p", NS)
    assert "⟦" not in para_text(paras[2])
    assert all(run_has_highlight(r) for r in paras[2].findall("w:r", NS))


def test_refuses_to_overwrite_input(built, tmp_path, capsys):
    src, results, _ = built
    victim = tmp_path / "manuscript.docx"
    shutil.copyfile(src, victim)
    before = sha256(victim)
    assert run_fill("--docx", victim, "--results", results, "--out", victim) == 2
    alias = tmp_path / "sub" / ".." / "manuscript.docx"
    (tmp_path / "sub").mkdir()
    assert run_fill("--docx", victim, "--results", results, "--out", alias) == 2
    assert "refusing to overwrite" in capsys.readouterr().err
    assert sha256(victim) == before


def test_missing_figure_file(built, tmp_path):
    src, _, _ = built
    results = tmp_path / "results"
    write_results(results)
    (results / "fig_selective.png").unlink()
    out = tmp_path / "out.docx"
    rj = tmp_path / "r.json"
    assert run_fill("--docx", src, "--results", results, "--out", out, "--report-json", rj) == 1
    report = json.loads(rj.read_text(encoding="utf-8"))
    assert len(report["missing_figures"]) == 1 and "fig_selective.png" in report["missing_figures"][0]
    _, blobs = read_zip(out)
    assert b"PLACEHOLDER_FIG6_SELECTIVE" in blobs["word/document.xml"]
    assert figure_info(blobs, "Figure 5")
    assert run_fill("--docx", src, "--results", results, "--out", out, "--allow-missing") == 0


def test_split_token_interrupted_by_tab_is_error(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = highlighted_run(p, "⟦T5_R2")
    run.add_tab()
    run.add_text("_ACC⟧")
    doc.add_paragraph("Fine: ⟦T5_R2_NLL⟧")
    src = tmp_path / "tab.docx"
    doc.save(str(src))
    results = tmp_path / "results"
    write_results(results)
    out = tmp_path / "out.docx"
    rj = tmp_path / "r.json"
    assert run_fill("--docx", src, "--results", results, "--out", out, "--report-json", rj,
                    "--allow-missing") == 1          # errors are fatal even with --allow-missing
    report = json.loads(rj.read_text(encoding="utf-8"))
    assert any("malformed" in e for e in report["errors"])
    assert report["replaced"] == {"T5_R2_NLL": 1}


def test_shared_jpeg_placeholder_gets_new_png_parts(tmp_path):
    """Both placeholders reuse one JPEG part (python-docx de-duplicates images)."""
    ph = tmp_path / "ph.jpg"
    Image.new("RGB", (2100, 780), (128, 128, 128)).save(ph, format="JPEG")
    doc = Document()
    doc.add_paragraph("Figures:")
    rename_docpr(doc.add_picture(str(ph), width=Inches(6.0)), FIG5)
    rename_docpr(doc.add_picture(str(ph), width=Inches(6.0)), FIG6)
    src = tmp_path / "shared.docx"
    doc.save(str(src))
    # drop any png Default so that the Content_Types fix-up is exercised
    infos, blobs = read_zip(src)
    ct = etree.fromstring(blobs["[Content_Types].xml"])
    for d in ct.findall("ct:Default", NS):
        if d.get("Extension").lower() == "png":
            ct.remove(d)
    blobs["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8",
                                                  standalone=True)
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as zf:
        for info in infos:
            zf.writestr(info.filename, blobs[info.filename])
    in5 = figure_info(blobs, FIG5)[0]
    assert in5["media"] == figure_info(blobs, FIG6)[0]["media"]
    old_media = in5["media"]
    assert old_media.endswith((".jpg", ".jpeg"))

    results = tmp_path / "results"
    pngs = write_results(results)
    out = tmp_path / "out.docx"
    assert run_fill("--docx", src, "--results", results, "--out", out) == 0
    _, out_blobs = read_zip(out)
    f5, f6 = figure_info(out_blobs, "Figure 5")[0], figure_info(out_blobs, "Figure 6")[0]
    assert f5["media"] != f6["media"]
    assert f5["media"].endswith(".png") and f6["media"].endswith(".png")
    assert out_blobs[f5["media"]] == pngs[FIG5] and out_blobs[f6["media"]] == pngs[FIG6]
    assert old_media not in out_blobs                        # orphaned JPEG removed
    assert old_media not in rels_targets(out_blobs, "word/document.xml").values()
    ct = etree.fromstring(out_blobs["[Content_Types].xml"])
    assert any(d.get("Extension") == "png" and d.get("ContentType") == "image/png"
               for d in ct.findall("ct:Default", NS))
    # 2100x780 placeholder -> 2100x840 figure: fig6 height recomputed, fig5 unchanged
    assert f5["cy"] == in5["cy"]
    assert f6["cy"] == round(f6["cx"] * 840 / 2100)
    assert len(Document(str(out)).inline_shapes) == 2


def test_load_fill_values_validation(tmp_path):
    path = tmp_path / "fill_values.json"
    path.write_text(json.dumps({
        "T5_R1_ACC": 0.8116,                     # not a string -> error
        "T6_R1_DACC": "-0.43 [-0.80, 0.05]",     # ASCII minus -> warning
        "TXT_PAIRED": "Line one.\nLine two.",     # newline collapsed
        "lower_key": "x",                        # can never match a token
    }), encoding="utf-8")
    report = fm.FillReport()
    values = fm.load_fill_values(path, report)
    assert "T5_R1_ACC" not in values and any("T5_R1_ACC" in e for e in report.errors)
    assert values["TXT_PAIRED"] == "Line one. Line two."
    assert values["T6_R1_DACC"] == "-0.43 [-0.80, 0.05]"
    assert any("U+2212" in w for w in report.warnings)
    assert "lower_key" not in values


def test_png_size_and_contract_keys():
    png = WORK / "assets" / "size_probe.png"
    png.parent.mkdir(parents=True, exist_ok=True)
    assert fm.png_size(make_png(png, (2100, 840), (0, 0, 0))) == (2100, 840)
    with pytest.raises(ValueError):
        fm.png_size(b"GIF89a" + b"\0" * 30)
    keys = fm.contract_keys()
    assert len(keys) == len(set(keys)) == 26 + 25 + 24 + 3 + 6


# ---------------------------------------------------------------------------
# regression tests from the adversarial review
# ---------------------------------------------------------------------------

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TRACK = 'w:author="A" w:date="2026-01-01T00:00:00Z"'


def _xml(fragment: str):
    from docx.oxml import parse_xml
    return parse_xml(fragment)


def _save(doc, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def _fill(src: Path, tmp_path: Path, *extra: str, values: dict | None = None,
          tag: str = "out") -> tuple[int, dict, Path]:
    results = tmp_path / f"results_{tag}"
    write_results(results, values)
    out, rj = tmp_path / f"{tag}.docx", tmp_path / f"{tag}.json"
    code = run_fill("--docx", src, "--results", results, "--out", out, "--report-json", rj,
                    *extra)
    return code, json.loads(rj.read_text(encoding="utf-8")), out


def test_load_fill_values_hostile_inputs(tmp_path):
    """BOM, duplicate keys, key with newline, whitespace, tab, empty, XML-illegal chars."""
    body = ('{"T5_R1_ACC": "first", "T5_R1_ACC": "second", "T5_R2_ACC\\n": "x", '
            '"T5_R3_ACC": "  0.8116 ", "T5_R4_ACC": "1.0\\t2.0", "T5_R5_ACC": "   ", '
            '"T5_TEMP": "1.2\\ufffe", "T6_R1_P": "0.1\\ud800", '
            '"T6_R2_P": "A & B < C > D \\u2212"}')
    path = tmp_path / "fill_values.json"
    path.write_bytes(b"\xef\xbb\xbf" + body.encode("utf-8"))        # UTF-8 BOM
    report = fm.FillReport()
    values = fm.load_fill_values(path, report)
    assert values["T5_R1_ACC"] == "second"
    assert any("duplicate key 'T5_R1_ACC'" in w for w in report.warnings)
    assert "T5_R2_ACC\n" not in values and "T5_R2_ACC" not in values
    assert values["T5_R3_ACC"] == "0.8116"
    assert values["T5_R4_ACC"] == "1.0 2.0"
    assert values["T6_R2_P"] == "A & B < C > D −"
    for key in ("T5_R5_ACC", "T5_TEMP", "T6_R1_P"):                   # errors, left unfilled
        assert key not in values and any(key in e for e in report.errors), key


def test_values_with_xml_specials_and_illegal_chars_do_not_crash(tmp_path):
    doc = Document()
    doc.add_paragraph("a=⟦T5_R1_ACC⟧ b=⟦T5_R2_ACC⟧ c=⟦T5_R3_ACC⟧")
    src = _save(doc, tmp_path / "in.docx")
    values = fill_values()
    values.update(T5_R1_ACC="]]></w:t><w:t>&amp;", T5_R2_ACC="−0.43 ⟨x⟩ 日本",
                  T5_R3_ACC="0.1￾")
    code, report, out = _fill(src, tmp_path, "--allow-missing", values=values)
    assert code == 1                                                   # error for T5_R3_ACC
    assert report["missing"] == {"T5_R3_ACC": 1}
    text = Document(str(out)).paragraphs[0].text
    assert text == "a=]]></w:t><w:t>&amp; b=−0.43 ⟨x⟩ 日本 c=⟦T5_R3_ACC⟧"


def test_deleted_and_moved_from_text_is_never_touched(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Kept ")
    p._p.append(_xml(f'<w:del xmlns:w="{W}" w:id="1" {TRACK}><w:r>'
                     '<w:delText>⟦T5_R1_ACC⟧</w:delText></w:r></w:del>'))
    p._p.append(_xml(f'<w:del xmlns:w="{W}" w:id="2" {TRACK}><w:r>'
                     '<w:delText>⟦ZZ_ONLY_DELETED⟧</w:delText></w:r></w:del>'))
    p._p.append(_xml(f'<w:moveFrom xmlns:w="{W}" w:id="3" {TRACK}><w:r>'
                     '<w:t>⟦T5_R2_ACC⟧</w:t></w:r></w:moveFrom>'))
    p._p.append(_xml(f'<w:moveTo xmlns:w="{W}" w:id="4" {TRACK}><w:r>'
                     '<w:t>⟦T5_R2_ACC⟧</w:t></w:r></w:moveTo>'))
    p._p.append(_xml(f'<w:ins xmlns:w="{W}" w:id="5" {TRACK}><w:r><w:rPr>'
                     '<w:highlight w:val="yellow"/></w:rPr><w:t>⟦T5_R3_ACC⟧</w:t></w:r></w:ins>'))
    src = _save(doc, tmp_path / "in.docx")
    code, report, out = _fill(src, tmp_path, "--allow-missing")
    assert code == 0
    assert report["replaced"] == {"T5_R2_ACC": 1, "T5_R3_ACC": 1}   # moveTo once, ins once
    assert report["missing"] == {} and report["errors"] == []
    _, blobs = read_zip(out)
    body = etree.fromstring(blobs["word/document.xml"])
    assert [t.text for t in body.iter(f"{{{W}}}delText")] == ["⟦T5_R1_ACC⟧",
                                                              "⟦ZZ_ONLY_DELETED⟧"]
    assert body.find(".//w:moveFrom/w:r/w:t", NS).text == "⟦T5_R2_ACC⟧"
    assert body.find(".//w:moveTo/w:r/w:t", NS).text == "v_t5_r2_acc"
    ins_run = body.find(".//w:ins/w:r", NS)
    assert ins_run.find("w:t", NS).text == fill_values()["T5_R3_ACC"]
    assert not run_has_highlight(ins_run)


def test_stray_notation_is_warning_but_token_like_fragments_are_errors(tmp_path):
    # The revision tool's own author comment says: text marked ⟦...⟧ is a placeholder.
    doc = Document()
    p = doc.add_paragraph("Fill ⟦T5_R1_ACC⟧.")
    doc.add_comment(p.runs, text="All text marked ⟦...⟧ is a placeholder.", author="R")
    code, report, _ = _fill(_save(doc, tmp_path / "ok_in.docx"), tmp_path, "--allow-missing",
                            tag="ok")
    assert code == 0 and report["errors"] == []
    assert any("comments.xml" in w and "not part of a token" in w for w in report["warnings"])

    # a malformed (unclosed) token must not swallow the following well-formed token
    doc = Document()
    doc.add_paragraph("Nested ⟦T5_R2_ACC: see ⟦T5_TEMP⟧ unclosed")
    doc.add_paragraph("Spaced ⟦ T5_R3_ACC ⟧")
    code, report, _ = _fill(_save(doc, tmp_path / "bad_in.docx"), tmp_path, "--allow-missing",
                            tag="bad")
    assert code == 1
    assert report["replaced"] == {"T5_TEMP": 1}
    assert len(report["errors"]) == 2 and all("malformed" in e for e in report["errors"])


def test_token_in_equation_is_reported_not_skipped(tmp_path):
    doc = Document()
    p = doc.add_paragraph("Eq: ")
    p._p.append(_xml(f'<m:oMath xmlns:m="{M}"><m:r><m:t>⟦T5_R1_ACC⟧</m:t></m:r></m:oMath>'))
    code, report, _ = _fill(_save(doc, tmp_path / "in.docx"), tmp_path, "--allow-missing")
    assert code == 1
    assert any("<m:t>" in e for e in report["errors"])


def test_textbox_choice_and_fallback_filled_but_counted_once(tmp_path):
    box = ('<w:txbxContent><w:p><w:r><w:t>⟦T5_R1_ACC⟧ ⟦T5_R2_</w:t></w:r>'
           '<w:r><w:t>ACC⟧</w:t></w:r></w:p></w:txbxContent>')
    doc = Document()
    doc.add_paragraph()._p.append(_xml(
        f'<w:r xmlns:w="{W}" xmlns:mc="{MC}"><mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:drawing>{box}</w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent></w:r>'))
    code, report, out = _fill(_save(doc, tmp_path / "in.docx"), tmp_path, "--allow-missing")
    assert code == 0
    assert report["replaced"] == {"T5_R1_ACC": 1, "T5_R2_ACC": 1}
    assert len(report["split_tokens"]) == 1
    _, blobs = read_zip(out)
    texts = [para_text(p) for p in etree.fromstring(blobs["word/document.xml"]).iter(f"{{{W}}}p")
             if p.getparent().tag == f"{{{W}}}txbxContent"]
    assert texts == ["v_t5_r1_acc v_t5_r2_acc"] * 2


def test_placeholder_sharing_media_with_ordinary_figure_and_header(tmp_path):
    ph = tmp_path / "ph.png"
    make_png(ph, (2100, 780), (128, 128, 128))
    doc = Document()
    doc.add_picture(str(ph), width=Inches(5))                    # ordinary figure, same rId
    rename_docpr(doc.add_picture(str(ph), width=Inches(6)), FIG5)
    hdr_run = doc.sections[0].header.paragraphs[0].add_run()
    rename_docpr(hdr_run.add_picture(str(ph), width=Inches(6)), FIG6)   # via header rels
    src = _save(doc, tmp_path / "in.docx")
    results = tmp_path / "results"
    pngs = write_results(results)
    out = tmp_path / "out.docx"
    assert run_fill("--docx", src, "--results", results, "--out", out) == 0
    _, in_blobs = read_zip(src)
    _, out_blobs = read_zip(out)
    old_media = figure_info(in_blobs, FIG5)[0]["media"]
    assert out_blobs[old_media] == in_blobs[old_media]               # ordinary figure intact
    doc_targets = rels_targets(out_blobs, "word/document.xml")
    blips = etree.fromstring(out_blobs["word/document.xml"]).findall(".//a:blip", NS)
    assert doc_targets[blips[0].get(f"{{{NS['r']}}}embed")] == old_media
    f5 = figure_info(out_blobs, "Figure 5")[0]
    assert f5["media"] != old_media and out_blobs[f5["media"]] == pngs[FIG5]
    hdr = next(n for n in out_blobs if n.startswith("word/header") and n.endswith(".xml"))
    blip = etree.fromstring(out_blobs[hdr]).find(".//a:blip", NS)
    target = rels_targets(out_blobs, hdr)[blip.get(f"{{{NS['r']}}}embed")]
    assert target != old_media and out_blobs[target] == pngs[FIG6]
    assert len(Document(str(out)).inline_shapes) == 2


def test_output_directory_is_usage_error(built, tmp_path, capsys):
    src, results, _ = built
    out_dir = tmp_path / "out.docx"
    out_dir.mkdir()
    assert run_fill("--docx", src, "--results", results, "--out", out_dir) == 2
    assert "not a directory" in capsys.readouterr().err


@pytest.mark.skipif(sys.platform != "win32", reason="Windows share-mode locking")
def test_output_locked_by_word_is_usage_error_without_temp_files(built, tmp_path, capsys):
    import ctypes
    from ctypes import wintypes
    src, results, _ = built
    locked = tmp_path / "locked.docx"
    shutil.copyfile(src, locked)
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = ctypes.c_void_p
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p]
    k32.CloseHandle.argtypes = [ctypes.c_void_p]
    # GENERIC_READ | GENERIC_WRITE with share mode 0, as Word holds an open document
    handle = k32.CreateFileW(str(locked), 0xC0000000, 0, None, 3, 0x80, None)
    assert handle not in (None, ctypes.c_void_p(-1).value)
    try:
        assert run_fill("--docx", src, "--results", results, "--out", locked) == 2
    finally:
        k32.CloseHandle(handle)
    assert "open in Word" in capsys.readouterr().err
    assert list(tmp_path.glob("*.tmp")) == []
