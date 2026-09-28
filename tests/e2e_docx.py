#!/usr/bin/env python3
"""Mock manuscript that holds EVERY CONTRACT.md section 4 placeholder, for end-to-end runs.

The document mimics the layout of the revised manuscript: Table 1 additions, Tables 5-7 with
one token per cell, the temperature and the six generated-prose tokens in their own
paragraphs, and the two figure placeholders (inline pictures whose ``<wp:docPr name=...>`` is
``PLACEHOLDER_FIG5_RELIABILITY`` / ``PLACEHOLDER_FIG6_SELECTIVE`` and whose PNGs have the
final pixel size). Every token run is yellow-highlighted; about half the tokens use the
``⟦KEY: description⟧`` form so both token spellings are exercised.

The key list is written out here from the contract text on purpose (not imported from the
scripts), so the end-to-end test can check both scripts against an independent copy.

Usage::

    python tests/e2e_docx.py --out _work/e2e/all_keys.docx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
PLACEHOLDER_DIR = ROOT / "tools" / "placeholders"

T5_ROWS = ("cb / argmax", "cb / PGS", "PGS / argmax (served)", "PGS / PGS",
           "PGS / argmax + temperature scaling")
T5_COLS = ("ACC", "CONF", "ECE", "BRIER", "NLL")
T6_ROWS = ("DINOv3-L vs. EVA-02-L", "DINOv3-L vs. DINOv2-L", "EVA-02-L vs. DINOv2-L",
           "PGS / PGS vs. PGS / argmax", "PGS / PGS vs. cb / argmax")
T6_COLS = ("DACC", "DF1", "BC", "P", "PHOLM")
T7_ROWS = ("1 − max p (served)", "1 − max p (PGS)", "Predictive entropy",
           "Expected entropy", "Mutual information", "Probability SD")
T7_COLS = ("AUROC", "AURC", "ACC80", "COV90")
T1_KEYS = ("T1_ECE", "T1_ENC", "T1_AUROC")
TXT_KEYS = ("TXT_REPRO", "TXT_CALIBRATION", "TXT_PAIRED", "TXT_SELECTIVE",
            "TXT_IMPLICATION", "TXT_CONCLUSION")
#: Text that precedes each prose token in its paragraph (TXT_REPRO continues a sentence).
TXT_LEAD = {"TXT_REPRO": "The exported records reproduce the accuracy and macro-F1 values of "
                         "Table 4 "}
FIGURES = {
    "PLACEHOLDER_FIG5_RELIABILITY": ("placeholder_fig5_reliability.png", (2100, 780)),
    "PLACEHOLDER_FIG6_SELECTIVE": ("placeholder_fig6_selective.png", (2100, 840)),
}


def contract_keys() -> list[str]:
    """CONTRACT.md section 4 keys, in contract order (independent literal copy)."""
    keys = [f"T5_R{r}_{c}" for r in range(1, 6) for c in T5_COLS] + ["T5_TEMP"]
    keys += [f"T6_R{r}_{c}" for r in range(1, 6) for c in T6_COLS]
    keys += [f"T7_R{r}_{c}" for r in range(1, 7) for c in T7_COLS]
    return keys + list(T1_KEYS) + list(TXT_KEYS)


def _token(key: str, with_description: bool) -> str:
    return f"⟦{key}: value for {key.lower()}⟧" if with_description else f"⟦{key}⟧"


def _add_token(paragraph, key: str, index: int) -> None:
    run = paragraph.add_run(_token(key, with_description=index % 2 == 1))
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _table(doc, caption: str, header: tuple[str, ...], rows: tuple[str, ...],
           prefix: str, cols: tuple[str, ...], counter: list[int]) -> None:
    doc.add_paragraph(caption)
    table = doc.add_table(rows=len(rows) + 1, cols=len(cols) + 1)
    table.style = "Table Grid"
    for j, h in enumerate(header):
        table.cell(0, j).text = h
    for i, label in enumerate(rows, start=1):
        table.cell(i, 0).text = label
        for j, col in enumerate(cols, start=1):
            _add_token(table.cell(i, j).paragraphs[0], f"{prefix}_R{i}_{col}", counter[0])
            counter[0] += 1


def _placeholder_png(name: str, size: tuple[int, int], assets: Path) -> Path:
    """The shipped placeholder PNG if it has the contract size, else a grey stand-in."""
    shipped = PLACEHOLDER_DIR / name
    if shipped.is_file():
        with Image.open(shipped) as im:
            if im.size == size:
                return shipped
    assets.mkdir(parents=True, exist_ok=True)
    path = assets / name
    Image.new("RGB", size, (235, 235, 235)).save(path, format="PNG")
    return path


def build_all_keys_docx(path: Path, assets: Path | None = None) -> list[str]:
    """Write the mock manuscript to ``path``; return the keys it contains (each once)."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    assets = assets or path.parent / "e2e_assets"
    doc = Document()
    counter = [0]
    doc.add_heading("End-to-end placeholder test manuscript", level=1)

    doc.add_paragraph("Table 1. Overview (additions).")
    t1 = doc.add_table(rows=len(T1_KEYS), cols=2)
    t1.style = "Table Grid"
    for i, (label, key) in enumerate(zip(("Calibration (ECE, served vs. PGS)",
                                          "Encoder comparison (DINOv3-L vs. EVA-02-L)",
                                          "Error detection (AUROC, 1 − max p vs. MI)"), T1_KEYS)):
        t1.cell(i, 0).text = label
        _add_token(t1.cell(i, 1).paragraphs[0], key, counter[0])
        counter[0] += 1

    _table(doc, "Table 5. Calibration (deployed pair, test set).",
           ("Configuration", "Accuracy", "Mean conf.", "ECE [95% CI]", "Brier", "NLL"),
           T5_ROWS, "T5", T5_COLS, counter)
    p = doc.add_paragraph("Fitted temperature T = ")
    _add_token(p, "T5_TEMP", counter[0])
    counter[0] += 1
    p.add_run(".")

    _table(doc, "Table 6. Paired tests (test set).",
           ("Comparison", "Δacc (pp) [95% CI]", "ΔF1 (pp) [95% CI]", "b / c", "p", "p (Holm)"),
           T6_ROWS, "T6", T6_COLS, counter)
    _table(doc, "Table 7. Selective prediction (deployed pair, PGS checkpoint).",
           ("Score", "AUROC [95% CI]", "AURC ×100", "Acc@80%", "Cov@90%"),
           T7_ROWS, "T7", T7_COLS, counter)

    for key in TXT_KEYS:
        p = doc.add_paragraph(TXT_LEAD.get(key, ""))
        _add_token(p, key, counter[0])
        counter[0] += 1

    for docpr_name, (png_name, size) in FIGURES.items():
        doc.add_paragraph(f"Figure placeholder {docpr_name}:")
        shape = doc.add_picture(str(_placeholder_png(png_name, size, assets)), width=Inches(6.5))
        shape._inline.docPr.set("name", docpr_name)

    doc.save(str(path))
    return contract_keys()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", type=Path, required=True, help="DOCX to write")
    args = parser.parse_args(argv)
    keys = build_all_keys_docx(args.out)
    print(f"wrote {args.out} with {len(keys)} tokens and {len(FIGURES)} figure placeholders")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
