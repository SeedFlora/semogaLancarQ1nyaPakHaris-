#!/usr/bin/env python3
"""Stage 2 of the per-sample uncertainty analysis (CONTRACT.md sections 3-5).

Reads the per-sample prediction export written by stage 1
(``export_per_sample_predictions.py``) and computes, on the TEST split, every quantity
defined in CONTRACT.md section 3.  The validation split is used for exactly one thing:
fitting the temperature of the served (``pgs/argmax``) probabilities.

Outputs (directory ``--out``)::

    results.json          every computed number, nested and self-describing
    fill_values.json      flat {KEY: formatted string} for every placeholder of section 4
    fig_reliability.png   7.0 x 2.6 in at 300 dpi = 2100 x 780 px (asserted after saving)
    fig_selective.png     7.0 x 2.8 in at 300 dpi = 2100 x 840 px (asserted after saving)
    tables.md             human-readable Tables 4-7 and the Table 1 additions
    tables/*.csv          the same tables (raw numbers) plus per-bin and curve data
    claims_check.md       verdicts for manuscript claims C1-C6 (section 5)

Statistical conventions (fixed by the contract):

* ECE: top-label, 15 equal-width bins (i/15, (i+1)/15], confidence = max p.
* Bootstrap: one index matrix ``default_rng(seed).integers(0, N, size=(B, N))`` drawn once and
  shared by every bootstrapped quantity; paired comparisons apply the same resample to both
  systems.  Intervals are 95 % percentile intervals (``numpy.percentile``, linear).
  The implementation is vectorised: the index matrix is turned into a (B, N) count matrix and
  every additive statistic is obtained as ``counts @ per_sample_statistic``.
* McNemar: exact two-sided binomial test on the discordant counts; Holm within each family.

Usage::

    python analyze_uncertainty.py --preds per_sample_predictions --out results \\
        [--deployed dinov3_large__mE5_large] [--bootstrap 2000] [--seed 42] \\
        [--manuscript SmartCitty_IJOST_Rev_1.docx]

``--manuscript`` is optional and only read (python-docx) to quote the current wording of the
claims in ``claims_check.md``; without it, the wording of the submitted revision is quoted.
"""

from __future__ import annotations

import argparse
import datetime as _dt
import functools
import json
import logging
import math
import platform
import re
import sys
import time
from dataclasses import dataclass, field
from fractions import Fraction
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

import numpy as np
import pandas as pd
from scipy import optimize, stats
from sklearn.metrics import f1_score, roc_auc_score

LOG = logging.getLogger("analyze_uncertainty")

# ---------------------------------------------------------------------------------------------
# Contract constants
# ---------------------------------------------------------------------------------------------

N_CLASSES = 9
N_BINS = 15
BIN_EDGES = np.linspace(0.0, 1.0, N_BINS + 1)
PROB_CLIP = 1e-12
TEMP_BOUNDS = (0.05, 20.0)
ALPHA = 0.05
DEFER_COVERAGE = 0.80                 # "accuracy at 80 % coverage"
TARGET_ACCURACY = Fraction(9, 10)     # "coverage at 90 % accuracy" (exact integer comparison)
DPI = 300
FIG_RELIABILITY_IN = (7.0, 2.6)
FIG_SELECTIVE_IN = (7.0, 2.8)
FIG_RELIABILITY_PX = (2100, 780)
FIG_SELECTIVE_PX = (2100, 840)
SUM_TOL = 1e-6
TIE_TOL = 1e-9

PROB_COLS = [f"p{k}" for k in range(N_CLASSES)]
Q_COLS = [f"q{k}" for k in range(N_CLASSES)]
BASE_COLUMNS = ["split", "row_id", "label", "pred", *PROB_COLS]
PGS_EXTRA_COLUMNS = ["mi", "pred_entropy", "exp_entropy", "prob_std", *Q_COLS, "pred_loglin"]
UNCERTAINTY_COLS = ["mi", "pred_entropy", "exp_entropy", "prob_std"]
FILE_RE = re.compile(
    r"^preds__(?P<image>.+?)__(?P<text>.+?)__(?P<checkpoint>cb|pgs)__(?P<mode>argmax|pgs)\.csv\.gz$"
)

DEFAULT_DEPLOYED = "dinov3_large__mE5_large"
DEFAULT_ENCODERS = ("dinov3_large", "eva02_large", "dinov2_large")

MINUS = "−"
ENDASH = "–"

DISPLAY_NAMES = {
    "dinov3_large": "DINOv3-L",
    "eva02_large": "EVA-02-L",
    "dinov2_large": "DINOv2-L",
    "hiera_large": "Hiera-L",
    "mE5_large": "mE5-L",
    "bge_m3": "BGE-M3",
    "indobert": "IndoBERT-L",
    "cendol_mt5": "Cendol-mT5-L",
}

#: Table 5 rows: (row, checkpoint, mode, temperature-scaled, description)
T5_ROWS: tuple[tuple[str, str, str, bool, str], ...] = (
    ("R1", "cb", "argmax", False, "Baseline checkpoint, argmax"),
    ("R2", "cb", "pgs", False, "Baseline checkpoint, PGS-averaged (M = 30)"),
    ("R3", "pgs", "argmax", False, "PGS checkpoint, argmax (served path)"),
    ("R4", "pgs", "pgs", False, "PGS checkpoint, PGS-averaged (M = 30)"),
    ("R5", "pgs", "argmax", True, "PGS checkpoint, argmax + temperature scaling"),
)

#: Table 7 rows: (row, score id, description, checkpoint/mode of the reference predictions)
T7_ROWS: tuple[tuple[str, str, str, str], ...] = (
    ("R1", "S1", f"1 {MINUS} max p, served argmax", "argmax"),
    ("R2", "S2", f"1 {MINUS} max p, PGS-averaged", "pgs"),
    ("R3", "S3", "Predictive entropy", "pgs"),
    ("R4", "S4", "Expected entropy (aleatoric)", "pgs"),
    ("R5", "S5", "Mutual information (epistemic)", "pgs"),
    ("R6", "S6", "Member-probability SD", "pgs"),
)
#: Phrases used in generated prose for each score (lower-case, mid-sentence).
SCORE_PROSE = {
    "S1": f"1 {MINUS} max p of the served argmax path",
    "S2": f"1 {MINUS} max p of the PGS-averaged probabilities",
    "S3": "the predictive entropy",
    "S4": "the expected (aleatoric) entropy",
    "S5": "the mutual information",
    "S6": "the standard deviation of the member probabilities",
}

COMBO_PROSE = {
    ("cb", "argmax"): "the baseline checkpoint with argmax inference",
    ("cb", "pgs"): "the baseline checkpoint with PGS-averaged inference",
    ("pgs", "argmax"): "the PGS checkpoint with argmax inference",
    ("pgs", "pgs"): "the PGS checkpoint with PGS-averaged inference",
}

# Okabe-Ito colour-blind-safe palette.
OI = {
    "black": "#000000",
    "orange": "#E69F00",
    "sky": "#56B4E9",
    "green": "#009E73",
    "yellow": "#F0E442",
    "blue": "#0072B2",
    "vermillion": "#D55E00",
    "purple": "#CC79A7",
    "grey": "#8C8C8C",
}


def _required_keys() -> list[str]:
    keys: list[str] = []
    for r in range(1, 6):
        keys += [f"T5_R{r}_{m}" for m in ("ACC", "CONF", "ECE", "BRIER", "NLL")]
    keys.append("T5_TEMP")
    for r in range(1, 6):
        keys += [f"T6_R{r}_{m}" for m in ("DACC", "DF1", "BC", "P", "PHOLM")]
    for r in range(1, 7):
        keys += [f"T7_R{r}_{m}" for m in ("AUROC", "AURC", "ACC80", "COV90")]
    keys += ["T1_ECE", "T1_ENC", "T1_AUROC"]
    keys += ["TXT_REPRO", "TXT_CALIBRATION", "TXT_PAIRED", "TXT_SELECTIVE",
             "TXT_IMPLICATION", "TXT_CONCLUSION"]
    return keys


#: Every placeholder key of CONTRACT.md section 4, in contract order.
REQUIRED_KEYS: list[str] = _required_keys()


class SchemaError(ValueError):
    """The stage-1 export violates the file contract (CONTRACT.md section 2)."""


# ---------------------------------------------------------------------------------------------
# Number formatting (CONTRACT.md section 4)
# ---------------------------------------------------------------------------------------------


def _finite(x: Any) -> bool:
    return x is not None and isinstance(x, (int, float, np.floating, np.integer)) and math.isfinite(float(x))


def fmt_fixed(x: float | None, dp: int, *, plus: bool = False) -> str:
    """Fixed-point with ``dp`` decimals, Unicode minus, optional explicit plus; never "-0.00"."""
    if not _finite(x):
        return "n/a"
    s = f"{float(x):.{dp}f}"
    if float(s) == 0.0:
        return f"{0.0:.{dp}f}"
    if plus and not s.startswith("-"):
        s = "+" + s
    return s.replace("-", MINUS)


def fmt_metric(x: float | None) -> str:
    """Metric with 4 decimals ("0.8116")."""
    return fmt_fixed(x, 4)


def fmt_ci(lo: float | None, hi: float | None, dp: int = 4) -> str:
    """Interval "[0.0101, 0.0150]"."""
    return f"[{fmt_fixed(lo, dp)}, {fmt_fixed(hi, dp)}]"


def fmt_metric_ci(x: float | None, lo: float | None, hi: float | None) -> str:
    """Metric with interval "0.0123 [0.0101, 0.0150]"."""
    return f"{fmt_metric(x)} {fmt_ci(lo, hi)}"


def fmt_pp(delta: float | None) -> str:
    """Difference of proportions in percentage points, 2 dp, explicit sign ("+0.12", "−0.43")."""
    return fmt_fixed(None if not _finite(delta) else 100.0 * float(delta), 2, plus=True)


def fmt_pp_ci(lo: float | None, hi: float | None) -> str:
    """Interval of a difference in pp, "[−0.30, 0.55]" (no plus signs inside)."""
    f = (lambda v: None if not _finite(v) else 100.0 * float(v))
    return fmt_ci(f(lo), f(hi), dp=2)


def fmt_p(p: float | None) -> str:
    """p-value: 3 significant digits ("0.0421", "0.237", "1.00"); "< 0.001" below 0.001."""
    if not _finite(p):
        return "n/a"
    p = float(p)
    if p < 0.001:
        return "< 0.001"
    exponent = math.floor(math.log10(p))
    dp = max(0, 2 - exponent)
    s = f"{p:.{dp}f}"
    if float(s) >= 10.0 ** (exponent + 1) and dp > 0:   # rounding carried into a new digit
        s = f"{p:.{dp - 1}f}"
    return s


def p_phrase(p: float | None, label: str = "p") -> str:
    """"p = 0.0421" or "p < 0.001"."""
    s = fmt_p(p)
    return f"{label} {s}" if s.startswith("<") else f"{label} = {s}"


def fmt_int(n: int) -> str:
    """Integer with thousands separators ("9,266")."""
    return f"{int(n):,}"


def fmt_temp(t: float | None) -> str:
    return fmt_fixed(t, 2)


def fmt_cov(c: float | None) -> str:
    return fmt_fixed(c, 3)


def fmt_aurc(a: float | None) -> str:
    """AURC reported x100 with 2 dp."""
    return fmt_fixed(None if not _finite(a) else 100.0 * float(a), 2)


def cmp_formatted(a: float, b: float, fmt: Callable[[float], str]) -> int:
    """Sign of a - b, but 0 when both print identically (keeps prose consistent with tables)."""
    if fmt(a) == fmt(b):
        return 0
    return 1 if a > b else -1


def join_list(items: Sequence[str]) -> str:
    """"A", "A and B", "A, B, and C"."""
    items = list(items)
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    if len(items) == 2:
        return f"{items[0]} and {items[1]}"
    return ", ".join(items[:-1]) + ", and " + items[-1]


def number_word(n: int) -> str:
    words = ["zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine"]
    return words[n] if 0 <= n < len(words) else fmt_int(n)


def display(name: str) -> str:
    return DISPLAY_NAMES.get(name, name)


# ---------------------------------------------------------------------------------------------
# Loading and validating the stage-1 export
# ---------------------------------------------------------------------------------------------


def expected_columns(mode: str) -> list[str]:
    return BASE_COLUMNS + (PGS_EXTRA_COLUMNS if mode == "pgs" else [])


@dataclass
class PredictionFile:
    """One ``preds__{img}__{txt}__{ckpt}__{mode}.csv.gz`` file."""

    image: str
    text: str
    checkpoint: str
    mode: str
    path: Path
    frame: pd.DataFrame
    manifest_entry: dict[str, Any]
    warnings: list[str] = field(default_factory=list)

    @property
    def pair(self) -> str:
        return f"{self.image}__{self.text}"

    @property
    def key(self) -> tuple[str, str, str]:
        return (self.pair, self.checkpoint, self.mode)


@dataclass
class View:
    """Arrays of one prediction file restricted to one split, in a fixed row order."""

    name: str
    row_id: np.ndarray
    label: np.ndarray
    pred: np.ndarray
    probs: np.ndarray
    extra: dict[str, np.ndarray]

    @property
    def n(self) -> int:
        return int(self.label.shape[0])

    @property
    def correct(self) -> np.ndarray:
        return self.pred == self.label

    @property
    def confidence(self) -> np.ndarray:
        return self.probs.max(axis=1)


def _argmax_tie_ok(probs: np.ndarray, pred: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Rows whose ``pred`` differs from argmax, and whether each is a written-precision tie."""
    amax = probs.argmax(axis=1)
    bad = np.flatnonzero(amax != pred)
    tie = probs[bad, pred[bad]] >= probs[bad].max(axis=1) - TIE_TOL
    return bad, tie


def validate_frame(df: pd.DataFrame, mode: str, name: str,
                   n_val: int | None = None, n_test: int | None = None) -> list[str]:
    """Check one export frame against CONTRACT.md section 2. Raise SchemaError; return warnings."""
    warnings: list[str] = []
    cols = expected_columns(mode)
    if list(df.columns) != cols:
        raise SchemaError(f"{name}: columns {list(df.columns)} differ from the contract {cols}")
    if len(df) == 0:
        raise SchemaError(f"{name}: empty file")
    split = df["split"].astype(str).to_numpy()
    bad_split = sorted(set(split) - {"val", "test"})
    if bad_split:
        raise SchemaError(f"{name}: unexpected split values {bad_split}")
    is_test = split == "test"
    if is_test.any():
        first_test = int(np.argmax(is_test))
        if (~is_test[first_test:]).any():
            raise SchemaError(f"{name}: validation rows must precede test rows")
    got_val, got_test = int((~is_test).sum()), int(is_test.sum())
    if n_val is not None and got_val != n_val:
        raise SchemaError(f"{name}: {got_val} val rows, manifest says n_val = {n_val}")
    if n_test is not None and got_test != n_test:
        raise SchemaError(f"{name}: {got_test} test rows, manifest says n_test = {n_test}")
    if got_test == 0:
        raise SchemaError(f"{name}: no test rows")
    int_cols = ["row_id", "label", "pred"] + (["pred_loglin"] if mode == "pgs" else [])
    for c in int_cols:
        if not pd.api.types.is_integer_dtype(df[c]):
            raise SchemaError(f"{name}: column {c!r} must be integer, got {df[c].dtype}")
    for c in ["label", "pred"] + (["pred_loglin"] if mode == "pgs" else []):
        v = df[c].to_numpy()
        if v.min() < 0 or v.max() >= N_CLASSES:
            raise SchemaError(f"{name}: column {c!r} outside 0..{N_CLASSES - 1}")
    if df.duplicated(["split", "row_id"]).any():
        raise SchemaError(f"{name}: duplicated (split, row_id) keys")

    def check_probs(block: list[str], pred_col: str) -> None:
        p = df[block].to_numpy(dtype=np.float64)
        if not np.isfinite(p).all():
            raise SchemaError(f"{name}: non-finite values in {block[0]}..{block[-1]}")
        if (p < 0).any() or (p > 1 + TIE_TOL).any():
            raise SchemaError(f"{name}: probabilities outside [0, 1] in {block[0]}..{block[-1]}")
        dev = np.abs(p.sum(axis=1) - 1.0)
        if (dev > SUM_TOL).any():
            raise SchemaError(f"{name}: {int((dev > SUM_TOL).sum())} rows of {block[0]}..{block[-1]} "
                              f"do not sum to 1 +- {SUM_TOL} (max deviation {dev.max():.3g})")
        bad, tie = _argmax_tie_ok(p, df[pred_col].to_numpy())
        if (~tie).any():
            raise SchemaError(f"{name}: {int((~tie).sum())} rows where {pred_col!r} is not the argmax "
                              f"of {block[0]}..{block[-1]}")
        if bad.size:
            warnings.append(f"{name}: {bad.size} rows where {pred_col!r} differs from argmax only "
                            "within written precision (tie); the exported value is kept")

    check_probs(PROB_COLS, "pred")
    if mode == "pgs":
        check_probs(Q_COLS, "pred_loglin")
        u = df[UNCERTAINTY_COLS].to_numpy(dtype=np.float64)
        if not np.isfinite(u).all():
            raise SchemaError(f"{name}: non-finite uncertainty values")
        if (u < -1e-9).any():
            raise SchemaError(f"{name}: negative uncertainty values")
        mi_re = np.maximum(df["pred_entropy"].to_numpy() - df["exp_entropy"].to_numpy(), 0.0)
        dev = np.abs(mi_re - df["mi"].to_numpy())
        if (dev > 1e-6).any():
            warnings.append(f"{name}: mi differs from max(pred_entropy - exp_entropy, 0) by up to "
                            f"{dev.max():.3g} on {int((dev > 1e-6).sum())} rows")
    return warnings


def load_prediction_dir(preds_dir: Path) -> tuple[dict[str, Any], dict[tuple[str, str, str], PredictionFile]]:
    """Load ``manifest.json`` and every prediction file listed in it (validated)."""
    manifest_path = preds_dir / "manifest.json"
    if not manifest_path.is_file():
        raise SchemaError(f"{manifest_path} not found")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for k in ("n_val", "n_test", "files"):
        if k not in manifest:
            raise SchemaError(f"manifest.json lacks {k!r}")
    n_val, n_test = int(manifest["n_val"]), int(manifest["n_test"])
    files: dict[tuple[str, str, str], PredictionFile] = {}
    listed = set()
    for entry in manifest["files"]:
        fname = entry["file"]
        listed.add(fname)
        m = FILE_RE.match(fname)
        if not m:
            raise SchemaError(f"manifest file name {fname!r} does not follow the contract pattern")
        parts = m.groupdict()
        for k_manifest, k_name in (("image", "image"), ("text", "text"),
                                   ("checkpoint", "checkpoint"), ("mode", "mode")):
            if k_manifest in entry and entry[k_manifest] != parts[k_name]:
                raise SchemaError(f"{fname}: manifest {k_manifest}={entry[k_manifest]!r} "
                                  f"contradicts the file name")
        path = preds_dir / fname
        if not path.is_file():
            raise SchemaError(f"{path} listed in manifest.json but missing")
        df = pd.read_csv(path, compression="gzip", dtype={"split": str})
        warns = validate_frame(df, parts["mode"], fname, n_val=n_val, n_test=n_test)
        pf = PredictionFile(image=parts["image"], text=parts["text"], checkpoint=parts["checkpoint"],
                            mode=parts["mode"], path=path, frame=df, manifest_entry=entry,
                            warnings=warns)
        if pf.key in files:
            raise SchemaError(f"duplicate manifest entry for {pf.key}")
        files[pf.key] = pf
    for extra in sorted(p.name for p in preds_dir.glob("preds__*.csv.gz")):
        if extra not in listed:
            LOG.warning("%s is present but not listed in manifest.json; ignored", extra)
    return manifest, files


def make_view(pf: PredictionFile, split: str, reference: pd.DataFrame | None = None) -> View:
    """Restrict a file to ``split``; if ``reference`` (split, row_id, label) is given, align to it."""
    df = pf.frame[pf.frame["split"] == split].reset_index(drop=True)
    name = f"{pf.path.name}[{split}]"
    if reference is not None:
        if len(df) != len(reference):
            raise SchemaError(f"{name}: {len(df)} rows vs {len(reference)} in the reference file")
        here = pd.MultiIndex.from_frame(df[["split", "row_id"]])
        want = pd.MultiIndex.from_frame(reference[["split", "row_id"]])
        idx = here.get_indexer(want)
        if (idx < 0).any():
            raise SchemaError(f"{name}: {int((idx < 0).sum())} (split, row_id) keys of the reference "
                              "file are missing")
        df = df.iloc[idx].reset_index(drop=True)
        if not np.array_equal(df["label"].to_numpy(), reference["label"].to_numpy()):
            raise SchemaError(f"{name}: labels differ from the reference file for the same row_id")
    extra = {c: df[c].to_numpy(dtype=np.float64) for c in UNCERTAINTY_COLS if c in df.columns}
    if "pred_loglin" in df.columns:
        extra["pred_loglin"] = df["pred_loglin"].to_numpy(dtype=np.int64)
    return View(name=name, row_id=df["row_id"].to_numpy(dtype=np.int64),
                label=df["label"].to_numpy(dtype=np.int64), pred=df["pred"].to_numpy(dtype=np.int64),
                probs=df[PROB_COLS].to_numpy(dtype=np.float64), extra=extra)


# ---------------------------------------------------------------------------------------------
# Core metrics
# ---------------------------------------------------------------------------------------------


def bin_index(conf: np.ndarray) -> np.ndarray:
    """Bin i = (i/15, (i+1)/15]; values are compared against the same edges as a naive loop."""
    return np.clip(np.searchsorted(BIN_EDGES, conf, side="left") - 1, 0, N_BINS - 1)


def reliability_table(conf: np.ndarray, correct: np.ndarray) -> dict[str, np.ndarray]:
    """Per-bin count, share, accuracy and mean confidence (NaN for empty bins)."""
    b = bin_index(conf)
    count = np.bincount(b, minlength=N_BINS).astype(np.int64)
    s_acc = np.bincount(b, weights=correct.astype(np.float64), minlength=N_BINS)
    s_conf = np.bincount(b, weights=conf, minlength=N_BINS)
    with np.errstate(invalid="ignore", divide="ignore"):
        acc = np.where(count > 0, s_acc / np.maximum(count, 1), np.nan)
        mconf = np.where(count > 0, s_conf / np.maximum(count, 1), np.nan)
    return {"lower": BIN_EDGES[:-1], "upper": BIN_EDGES[1:], "count": count,
            "share": count / conf.shape[0], "accuracy": acc, "mean_confidence": mconf}


def ece_score(probs: np.ndarray, labels: np.ndarray, pred: np.ndarray | None = None) -> float:
    """Top-label ECE with 15 equal-width bins."""
    conf = probs.max(axis=1)
    pred = probs.argmax(axis=1) if pred is None else pred
    correct = (pred == labels).astype(np.float64)
    b = bin_index(conf)
    gap = np.bincount(b, weights=correct - conf, minlength=N_BINS)
    return float(np.abs(gap).sum() / conf.shape[0])


def brier_per_sample(probs: np.ndarray, labels: np.ndarray) -> np.ndarray:
    onehot = np.zeros_like(probs)
    onehot[np.arange(labels.shape[0]), labels] = 1.0
    return ((probs - onehot) ** 2).sum(axis=1)


def nll_per_sample(probs: np.ndarray, labels: np.ndarray) -> np.ndarray:
    return -np.log(np.clip(probs[np.arange(labels.shape[0]), labels], PROB_CLIP, 1.0))


def macro_f1_from_confusion(cm: np.ndarray) -> np.ndarray:
    """Macro-F1 from confusion counts of shape (..., 9, 9) [true, pred].

    Equals sklearn ``f1_score(average='macro', labels=range(9), zero_division=0)``:
    F1_k = 2 TP / (2 TP + FP + FN), and 0 when the denominator is 0.
    """
    cm = np.asarray(cm, dtype=np.float64)
    tp = np.diagonal(cm, axis1=-2, axis2=-1)
    fp = cm.sum(axis=-2) - tp
    fn = cm.sum(axis=-1) - tp
    den = 2 * tp + fp + fn
    with np.errstate(invalid="ignore", divide="ignore"):
        f1 = np.where(den > 0, 2 * tp / np.where(den > 0, den, 1.0), 0.0)
    return f1.mean(axis=-1)


def confusion_onehot(labels: np.ndarray, pred: np.ndarray) -> np.ndarray:
    """(N, 81) indicator of the confusion cell label*9 + pred."""
    out = np.zeros((labels.shape[0], N_CLASSES * N_CLASSES))
    out[np.arange(labels.shape[0]), labels * N_CLASSES + pred] = 1.0
    return out


def macro_f1(labels: np.ndarray, pred: np.ndarray) -> float:
    return float(f1_score(labels, pred, average="macro", labels=list(range(N_CLASSES)), zero_division=0))


def mcnemar_exact(correct_a: np.ndarray, correct_b: np.ndarray) -> tuple[int, int, float]:
    """b = #(A right, B wrong), c = #(A wrong, B right), exact two-sided binomial p."""
    a = np.asarray(correct_a, dtype=bool)
    bb = np.asarray(correct_b, dtype=bool)
    b = int(np.sum(a & ~bb))
    c = int(np.sum(~a & bb))
    if b + c == 0:
        return b, c, 1.0
    return b, c, float(stats.binomtest(min(b, c), b + c, 0.5).pvalue)


def holm(pvalues: Sequence[float]) -> list[float]:
    """Holm step-down adjusted p-values (monotone, capped at 1)."""
    p = np.asarray(pvalues, dtype=np.float64)
    m = p.shape[0]
    order = np.argsort(p, kind="stable")
    adjusted = np.empty(m)
    running = 0.0
    for rank, i in enumerate(order):
        running = max(running, min(1.0, (m - rank) * p[i]))
        adjusted[i] = running
    return adjusted.tolist()


def auroc_from_weights(weights: np.ndarray, score: np.ndarray, positive: np.ndarray) -> np.ndarray:
    """AUROC (ties count 1/2) for each row of non-negative sample weights ``(R, N)``.

    With unit weights this is ``roc_auc_score(positive, score)``; with bootstrap counts it is the
    AUROC of the resampled data set.  NaN where a row has no positive or no negative weight.
    """
    weights = np.atleast_2d(np.asarray(weights, dtype=np.float64))
    order = np.argsort(score, kind="stable")
    s = np.asarray(score, dtype=np.float64)[order]
    pos = np.asarray(positive, dtype=bool)[order]
    new_group = np.r_[True, s[1:] != s[:-1]]
    group = np.cumsum(new_group) - 1
    starts = np.flatnonzero(new_group)
    ends = np.r_[starts[1:], s.shape[0]]
    pos_cols = np.flatnonzero(pos)
    gs, ge = starts[group[pos_cols]], ends[group[pos_cols]]
    w = weights[:, order]
    wneg = np.where(pos, 0.0, w)
    cneg = np.concatenate([np.zeros((w.shape[0], 1)), np.cumsum(wneg, axis=1)], axis=1)
    below = cneg[:, gs]
    equal = cneg[:, ge] - below
    wpos = w[:, pos_cols]
    num = (wpos * (below + 0.5 * equal)).sum(axis=1)
    den = wpos.sum(axis=1) * wneg.sum(axis=1)
    with np.errstate(invalid="ignore", divide="ignore"):
        return np.where(den > 0, num / np.where(den > 0, den, 1.0), np.nan)


def auroc(score: np.ndarray, incorrect: np.ndarray) -> float:
    """``roc_auc_score(incorrect, score)``; NaN when only one class is present."""
    incorrect = np.asarray(incorrect, dtype=bool)
    if incorrect.all() or not incorrect.any():
        return float("nan")
    return float(roc_auc_score(incorrect.astype(int), score))


@dataclass
class SelectiveResult:
    aurc: float
    acc_at_coverage: float
    k_at_coverage: int
    coverage_at_accuracy: float
    k_at_accuracy: int
    coverage: np.ndarray
    accuracy: np.ndarray


def selective_metrics(score: np.ndarray, correct: np.ndarray,
                      coverage_target: float = DEFER_COVERAGE,
                      accuracy_target: Fraction = TARGET_ACCURACY) -> SelectiveResult:
    """Accuracy-coverage analysis; samples accepted in ascending score order (stable)."""
    n = score.shape[0]
    order = np.argsort(score, kind="stable")
    cum = np.cumsum(np.asarray(correct, dtype=np.int64)[order])
    k = np.arange(1, n + 1)
    acc = cum / k
    aurc = float(np.mean((k - cum) / k))
    k_cov = int(round(coverage_target * n))
    k_cov = min(max(k_cov, 1), n)
    ok = cum * accuracy_target.denominator >= accuracy_target.numerator * k
    k_acc = int(k[ok].max()) if ok.any() else 0
    return SelectiveResult(aurc=aurc, acc_at_coverage=float(cum[k_cov - 1] / k_cov), k_at_coverage=k_cov,
                           coverage_at_accuracy=k_acc / n, k_at_accuracy=k_acc,
                           coverage=k / n, accuracy=acc)


def oracle_curve(correct: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Accuracy-coverage curve of an oracle that accepts all correct predictions first."""
    n = correct.shape[0]
    n_correct = int(np.sum(correct))
    k = np.arange(1, n + 1)
    return k / n, np.minimum(k, n_correct) / k


def softmax(z: np.ndarray, axis: int = -1) -> np.ndarray:
    z = z - z.max(axis=axis, keepdims=True)
    e = np.exp(z)
    return e / e.sum(axis=axis, keepdims=True)


def temperature_nll(t: float, logp: np.ndarray, labels: np.ndarray) -> float:
    z = logp / t
    z = z - z.max(axis=1, keepdims=True)
    lse = np.log(np.exp(z).sum(axis=1))
    return float(np.mean(lse - z[np.arange(labels.shape[0]), labels]))


def fit_temperature(probs: np.ndarray, labels: np.ndarray,
                    bounds: tuple[float, float] = TEMP_BOUNDS) -> dict[str, float]:
    """T minimising the NLL of softmax(log clip(p) / T) (bounded scalar search)."""
    logp = np.log(np.clip(probs, PROB_CLIP, 1.0))
    res = optimize.minimize_scalar(temperature_nll, bounds=bounds, method="bounded",
                                   args=(logp, labels), options={"xatol": 1e-7, "maxiter": 500})
    t = float(res.x)
    at_bound = min(abs(t - bounds[0]), abs(t - bounds[1])) < 1e-4
    if at_bound:
        LOG.warning("temperature %.4f lies at the search bound %s", t, bounds)
    return {"temperature": t, "val_nll_before": temperature_nll(1.0, logp, labels),
            "val_nll_after": float(res.fun), "at_bound": bool(at_bound), "converged": bool(res.success)}


def apply_temperature(probs: np.ndarray, t: float) -> np.ndarray:
    return softmax(np.log(np.clip(probs, PROB_CLIP, 1.0)) / t, axis=1)


# ---------------------------------------------------------------------------------------------
# Vectorised bootstrap
# ---------------------------------------------------------------------------------------------


class Bootstrap:
    """Fixed bootstrap resamples of ``n`` indices, stored as a (B, n) count matrix.

    The index matrix is ``numpy.random.default_rng(seed).integers(0, n, size=(B, n))``; row r of
    the count matrix holds how often each original index appears in resample r.  Any statistic
    that is a sum over samples is then ``counts @ per_sample_values`` (done in chunks).
    """

    def __init__(self, n: int, n_resamples: int = 2000, seed: int = 42, chunk: int = 250) -> None:
        self.n, self.n_resamples, self.seed, self.chunk = n, n_resamples, seed, chunk
        rng = np.random.default_rng(seed)
        idx = rng.integers(0, n, size=(n_resamples, n))
        dtype = np.uint16 if n < np.iinfo(np.uint16).max else np.uint32
        self.counts = np.empty((n_resamples, n), dtype=dtype)
        for s in range(0, n_resamples, chunk):
            rows = idx[s:s + chunk]
            c = rows.shape[0]
            flat = (rows + (np.arange(c, dtype=np.int64)[:, None] * n)).ravel()
            self.counts[s:s + c] = np.bincount(flat, minlength=c * n).reshape(c, n)
        self.first_indices = idx[:5].copy()   # kept for audits/tests

    def chunks(self) -> Iterable[np.ndarray]:
        for s in range(0, self.n_resamples, self.chunk):
            yield self.counts[s:s + self.chunk].astype(np.float64)

    def sums(self, values: np.ndarray) -> np.ndarray:
        """Resampled sums of per-sample ``values`` (N,) or (N, k) -> (B,) or (B, k)."""
        v = np.asarray(values, dtype=np.float64)
        flat = v.reshape(self.n, -1)
        out = np.concatenate([w @ flat for w in self.chunks()], axis=0)
        return out.reshape((self.n_resamples,) + v.shape[1:])

    def map(self, fn: Callable[[np.ndarray], np.ndarray]) -> np.ndarray:
        """Apply ``fn(weights_chunk) -> (chunk,)`` to all resamples."""
        return np.concatenate([fn(w) for w in self.chunks()], axis=0)

    @staticmethod
    def interval(values: np.ndarray, level: float = 0.95) -> tuple[float, float]:
        a = (1.0 - level) / 2.0 * 100.0
        lo, hi = np.nanpercentile(np.asarray(values, dtype=np.float64), [a, 100.0 - a])
        return float(lo), float(hi)


def ci_excludes_zero(lo: float, hi: float) -> bool:
    return lo > 0.0 or hi < 0.0


def intervals_overlap(a: Sequence[float], b: Sequence[float]) -> bool:
    return a[0] <= b[1] and b[0] <= a[1]


# ---------------------------------------------------------------------------------------------
# Analyses
# ---------------------------------------------------------------------------------------------


def calibration_row(probs: np.ndarray, labels: np.ndarray, pred: np.ndarray, boot: Bootstrap) -> dict[str, Any]:
    """Accuracy, mean confidence, ECE (+CI), Brier, NLL (+CIs) and the reliability bins."""
    n = labels.shape[0]
    conf = probs.max(axis=1)
    correct = (pred == labels).astype(np.float64)
    brier = brier_per_sample(probs, labels)
    nll = nll_per_sample(probs, labels)
    gap = np.zeros((n, N_BINS))
    gap[np.arange(n), bin_index(conf)] = correct - conf
    stacked = np.column_stack([correct, conf, brier, nll, gap])
    sums = boot.sums(stacked) / n
    ece_b = np.abs(sums[:, 4:]).sum(axis=1)
    point = stacked.mean(axis=0)
    ece = float(np.abs(point[4:]).sum())
    bins = reliability_table(conf, correct.astype(bool))
    return {
        "n": n,
        "accuracy": float(point[0]), "accuracy_ci": boot.interval(sums[:, 0]),
        "mean_confidence": float(point[1]), "mean_confidence_ci": boot.interval(sums[:, 1]),
        "ece": ece, "ece_ci": boot.interval(ece_b),
        "brier": float(point[2]), "brier_ci": boot.interval(sums[:, 2]),
        "nll": float(point[3]), "nll_ci": boot.interval(sums[:, 3]),
        "confidence_minus_accuracy": float(point[1] - point[0]),
        "bins": {k: np.asarray(v).tolist() for k, v in bins.items()},
    }


def analyse_calibration(views: Mapping[tuple[str, str], View], val_served: View,
                        boot: Bootstrap) -> dict[str, Any]:
    """Table 5 (deployed pair): four checkpoint/mode rows plus temperature-scaled served path."""
    temp = fit_temperature(val_served.probs, val_served.label)
    t = temp["temperature"]
    served = views[("pgs", "argmax")]
    ts_probs = apply_temperature(served.probs, t)
    ts_pred = ts_probs.argmax(axis=1)
    n_flip = int(np.sum(ts_pred != served.pred))
    if n_flip:
        LOG.warning("temperature scaling changed %d argmax decisions (written-precision ties); "
                    "the exported served predictions are used", n_flip)
    val_ts = apply_temperature(val_served.probs, t)
    temp.update({
        "fitted_on": val_served.name, "n_val": val_served.n,
        "val_ece_before": ece_score(val_served.probs, val_served.label, val_served.pred),
        "val_ece_after": ece_score(val_ts, val_served.label, val_served.pred),
        "test_argmax_changes": n_flip,
    })
    rows = []
    for row, ckpt, mode, scaled, desc in T5_ROWS:
        v = views[(ckpt, mode)]
        probs = ts_probs if scaled else v.probs
        r = calibration_row(probs, v.label, v.pred, boot)
        r.update({"row": row, "checkpoint": ckpt, "mode": mode, "temperature_scaled": scaled,
                  "description": desc, "source": v.name})
        rows.append(r)
    return {"rows": rows, "temperature": temp}


def paired_comparison(a: View, b: View, boot: Bootstrap) -> dict[str, Any]:
    """McNemar exact test and paired bootstrap of accuracy / macro-F1 differences (A - B)."""
    if not np.array_equal(a.label, b.label):
        raise SchemaError(f"paired comparison {a.name} vs {b.name}: labels are not aligned")
    bb, cc, p = mcnemar_exact(a.correct, b.correct)
    sums = boot.sums(np.hstack([confusion_onehot(a.label, a.pred), confusion_onehot(b.label, b.pred)]))
    cm_a = sums[:, :81].reshape(-1, N_CLASSES, N_CLASSES)
    cm_b = sums[:, 81:].reshape(-1, N_CLASSES, N_CLASSES)
    n = a.n
    acc_a_b = np.trace(cm_a, axis1=1, axis2=2) / n
    acc_b_b = np.trace(cm_b, axis1=1, axis2=2) / n
    f1_a_b, f1_b_b = macro_f1_from_confusion(cm_a), macro_f1_from_confusion(cm_b)
    acc_a, acc_b = float(a.correct.mean()), float(b.correct.mean())
    f1_a, f1_b = macro_f1(a.label, a.pred), macro_f1(b.label, b.pred)
    return {
        "source_a": a.name, "source_b": b.name, "n": n,
        "accuracy_a": acc_a, "accuracy_b": acc_b, "d_accuracy": acc_a - acc_b,
        "d_accuracy_ci": boot.interval(acc_a_b - acc_b_b),
        "macro_f1_a": f1_a, "macro_f1_b": f1_b, "d_macro_f1": f1_a - f1_b,
        "d_macro_f1_ci": boot.interval(f1_a_b - f1_b_b),
        "b": bb, "c": cc, "discordant": bb + cc, "p_mcnemar": p,
        "top1_disagreements": int(np.sum(a.pred != b.pred)),
    }


def analyse_paired(test_views: Mapping[tuple[str, str, str], View], deployed: str,
                   encoder_pairs: Sequence[str], boot: Bootstrap) -> dict[str, Any]:
    """Table 6 rows R1-R5 plus a supplementary baseline-checkpoint comparison."""
    e0, e1, e2 = encoder_pairs
    spec = [
        ("R1", "encoders", (e0, "pgs", "pgs"), (e1, "pgs", "pgs")),
        ("R2", "encoders", (e0, "pgs", "pgs"), (e2, "pgs", "pgs")),
        ("R3", "encoders", (e1, "pgs", "pgs"), (e2, "pgs", "pgs")),
        ("R4", "inference_modes", (deployed, "pgs", "pgs"), (deployed, "pgs", "argmax")),
        ("R5", "inference_modes", (deployed, "pgs", "pgs"), (deployed, "cb", "argmax")),
    ]
    rows = []
    for row, family, ka, kb in spec:
        r = paired_comparison(test_views[ka], test_views[kb], boot)
        r.update({"row": row, "family": family, "key_a": list(ka), "key_b": list(kb),
                  "label_a": _system_label(ka, deployed), "label_b": _system_label(kb, deployed)})
        rows.append(r)
    for family in ("encoders", "inference_modes"):
        members = [r for r in rows if r["family"] == family]
        for r, ph in zip(members, holm([r["p_mcnemar"] for r in members])):
            r["p_holm"] = ph
    supp = paired_comparison(test_views[(deployed, "cb", "pgs")], test_views[(deployed, "cb", "argmax")], boot)
    supp.update({"row": "S-cb", "family": "supplementary (not Holm-adjusted)",
                 "key_a": [deployed, "cb", "pgs"], "key_b": [deployed, "cb", "argmax"],
                 "label_a": "Baseline checkpoint, PGS-averaged", "label_b": "Baseline checkpoint, argmax"})
    return {"rows": rows, "supplementary": [supp],
            "families": {"encoders": ["R1", "R2", "R3"], "inference_modes": ["R4", "R5"]}}


def _system_label(key: Sequence[str], deployed: str) -> str:
    pair, ckpt, mode = key
    img, txt = pair.split("__", 1)
    enc = f"{display(img)} + {display(txt)}"
    ck = "PGS checkpoint" if ckpt == "pgs" else "baseline checkpoint"
    md = "PGS-averaged" if mode == "pgs" else "argmax"
    return f"{enc}, {ck}, {md}"


def analyse_selective(served: View, pgs: View, boot: Bootstrap) -> dict[str, Any]:
    """Table 7 (S1-S6), curves for the figure and paired AUROC differences."""
    scores = {
        "S1": 1.0 - served.probs.max(axis=1),
        "S2": 1.0 - pgs.probs.max(axis=1),
        "S3": pgs.extra["pred_entropy"],
        "S4": pgs.extra["exp_entropy"],
        "S5": pgs.extra["mi"],
        "S6": pgs.extra["prob_std"],
    }
    incorrect = {"argmax": ~served.correct, "pgs": ~pgs.correct}
    rows, boots, curves = [], {}, {}
    for row, sid, desc, ref in T7_ROWS:
        s, inc = scores[sid], incorrect[ref]
        point = auroc(s, inc)
        check = float(auroc_from_weights(np.ones((1, s.shape[0])), s, inc)[0])
        if np.isfinite(point) and abs(check - point) > 1e-9:
            raise AssertionError(f"vectorised AUROC {check} != sklearn {point} for {sid}")
        b = boot.map(lambda w, s=s, inc=inc: auroc_from_weights(w, s, inc))
        boots[sid] = b
        sel = selective_metrics(s, ~inc)
        curves[sid] = sel
        rows.append({
            "row": row, "score": sid, "description": desc,
            "reference_predictions": "pgs/argmax" if ref == "argmax" else "pgs/pgs",
            "auroc": point, "auroc_ci": boot.interval(b),
            "aurc": sel.aurc, "aurc_x100": 100.0 * sel.aurc,
            "acc_at_80": sel.acc_at_coverage, "k_at_80": sel.k_at_coverage,
            "coverage_at_90": sel.coverage_at_accuracy, "k_at_90": sel.k_at_accuracy,
            "full_coverage_accuracy": float(np.mean(~inc)),
            "n_errors": int(inc.sum()),
            "score_mean_correct": float(np.mean(s[~inc])) if (~inc).any() else None,
            "score_mean_incorrect": float(np.mean(s[inc])) if inc.any() else None,
        })
    diffs = []
    for a, b in (("S5", "S1"), ("S5", "S2"), ("S2", "S1")):
        d = boots[a] - boots[b]
        pa = next(r["auroc"] for r in rows if r["score"] == a)
        pb = next(r["auroc"] for r in rows if r["score"] == b)
        diffs.append({"a": a, "b": b, "d_auroc": pa - pb, "d_auroc_ci": boot.interval(d)})
    grid = np.round(np.linspace(0.01, 1.0, 100), 2)
    oc, oa = oracle_curve(pgs.correct)
    curve_table = {"coverage": grid.tolist()}
    for sid, sel in curves.items():
        k = np.clip(np.round(grid * sel.coverage.shape[0]).astype(int), 1, sel.coverage.shape[0])
        curve_table[sid] = sel.accuracy[k - 1].tolist()
    k = np.clip(np.round(grid * oc.shape[0]).astype(int), 1, oc.shape[0])
    curve_table["oracle_pgs"] = oa[k - 1].tolist()
    mi = scores["S5"]
    return {
        "rows": rows, "paired_auroc_differences": diffs, "curves_on_grid": curve_table,
        "mi_summary": {
            "n_zero": int(np.sum(mi <= 0)),
            "min_positive": float(mi[mi > 0].min()) if (mi > 0).any() else None,
            "median_correct": float(np.median(mi[pgs.correct])) if pgs.correct.any() else None,
            "median_incorrect": float(np.median(mi[~pgs.correct])) if (~pgs.correct).any() else None,
            "mean_prob_std": float(np.mean(scores["S6"])),
            "mean_mi": float(np.mean(mi)),
        },
        "_curves": curves, "_scores": scores, "_incorrect": incorrect,
    }


def analyse_reproduction(manifest: Mapping[str, Any], files: Mapping[tuple[str, str, str], PredictionFile],
                         test_views: Mapping[tuple[str, str, str], View], deployed: str) -> dict[str, Any]:
    """Recompute accuracy / macro-F1 per file and compare with the manuscript expectations."""
    rows = []
    for key, pf in sorted(files.items()):
        v = test_views[key]
        acc, f1 = float(v.correct.mean()), macro_f1(v.label, v.pred)
        entry = pf.manifest_entry
        exp = entry.get("expected_test") or {}
        checks = {}
        for metric, value in (("accuracy", acc), ("macro_f1", f1)):
            allowed = exp.get(metric)
            if allowed:
                allowed_s = [f"{float(x):.4f}" for x in allowed]
                nearest = min((float(x) for x in allowed), key=lambda x: abs(x - value))
                checks[metric] = {"value": value, "expected": [float(x) for x in allowed],
                                  "matches": f"{value:.4f}" in allowed_s,
                                  "deviation_pp": 100.0 * (value - nearest)}
        manifest_metrics = ((entry.get("metrics") or {}).get("test") or {})
        manifest_dev = {m: abs(float(manifest_metrics[m]) - val)
                        for m, val in (("accuracy", acc), ("macro_f1", f1)) if m in manifest_metrics}
        for m, dev in manifest_dev.items():
            if dev > 5e-5:
                LOG.warning("%s: recomputed test %s differs from manifest by %.2g", pf.path.name, m, dev)
        computed_flag = None if not checks else all(c["matches"] for c in checks.values())
        extra: dict[str, Any] = {}
        if pf.mode == "pgs" and "pred_loglin" in v.extra:
            extra = {"accuracy_loglin": float(np.mean(v.extra["pred_loglin"] == v.label)),
                     "macro_f1_loglin": macro_f1(v.label, v.extra["pred_loglin"]),
                     "mean_prob_std": float(np.mean(v.extra["prob_std"])),
                     "mean_mi": float(np.mean(v.extra["mi"]))}
        rows.append({
            "file": pf.path.name, "pair": pf.pair, "checkpoint": pf.checkpoint, "mode": pf.mode,
            "deployed_pair": pf.pair == deployed,
            "accuracy": acc, "macro_f1": f1, "checks": checks,
            "reproduces_computed": computed_flag,
            "reproduces_manifest": entry.get("reproduces_manuscript"),
            "manifest_metric_abs_diff": manifest_dev,
            "tree_count": entry.get("tree_count"), "cbm_sha256": entry.get("cbm_sha256"),
            "warnings": pf.warnings, **extra,
        })
    served = test_views[(deployed, "pgs", "argmax")]
    pgs = test_views[(deployed, "pgs", "pgs")]
    disagreements = int(np.sum(served.pred != pgs.pred))
    return {
        "rows": rows,
        "top1_disagreements_pgs_vs_argmax": disagreements,
        "top1_disagreements_manifest": manifest.get("top1_disagreements_pgs_vs_argmax"),
        "top1_disagreements_expected": 130,
        "manifest_meta": {k: manifest.get(k) for k in ("created_utc", "catboost_version", "numpy_version",
                                                       "virtual_ensembles", "pooling", "n_val", "n_test")},
    }


# ---------------------------------------------------------------------------------------------
# Claims (CONTRACT.md section 5)
# ---------------------------------------------------------------------------------------------

#: Wording of the submitted revision (SmartCitty_IJOST_Rev_1.docx), used when --manuscript is not given.
CLAIM_SOURCES: dict[str, dict[str, Any]] = {
    "C1": {
        "title": "PGS averaging lowers accuracy at a fixed tree count",
        "sections": "Sections 3.3, 4; Table 1",
        "phrases": ["lowers performance on both checkpoints", "PGS inference lowers",
                    "PGS effect at fixed trees"],
        "quotes": [
            "Once the tree count is held fixed, enabling PGS inference lowers performance on both "
            "checkpoints, by 0.43–0.82 pp in accuracy and 0.46–0.84 pp in macro-F1, while the whole "
            "apparent improvement is explained by the additional boosting iterations (+1.20 pp accuracy "
            "and +1.09 pp macro-F1 under the same inference mode).",
            "Methodologically, the study showed that the apparent accuracy benefit of Posterior Gaussian "
            "Sampling disappears once the number of boosting iterations is held constant: at a fixed tree "
            "count, PGS inference lowers macro-F1 by 0.46–0.84 pp, and the naive +0.63 pp difference is not "
            "distinguishable from sampling variance.",
            "Table 1: PGS effect at fixed trees | −0.46 to −0.84 pp macro-F1",
        ],
    },
    "C2": {
        "title": "The naive +0.63 pp macro-F1 difference cannot be distinguished from sampling variance",
        "sections": "Sections 3.3, 4, 5.1; Table 10",
        "phrases": ["cannot be distinguished from sampling variance", "not distinguishable from sampling variance",
                    "gain is not significant"],
        "quotes": [
            "The bootstrap analysis points the same way: the baseline macro-F1 (0.7684) falls inside the "
            "95% interval of the PGS configuration ([0.764, 0.784]), so even the naive +0.63 pp difference "
            "cannot be distinguished from sampling variance.",
            "Methodologically, the study showed that the apparent accuracy benefit of Posterior Gaussian "
            "Sampling disappears once the number of boosting iterations is held constant: at a fixed tree "
            "count, PGS inference lowers macro-F1 by 0.46–0.84 pp, and the naive +0.63 pp difference is not "
            "distinguishable from sampling variance.",
            "Table 10: Shows the naive +0.63 pp gain is not significant and PGS costs 0.46–0.84 pp at fixed "
            "trees; PGS is repositioned as a low-cost uncertainty signal (≈0.05 ms per sample) (Table 4)",
        ],
    },
    "C3": {
        "title": "DINOv3-L and EVA-02-L are indistinguishable on accuracy",
        "sections": "Sections 3.2, 5.1",
        "phrases": ["indistinguishable on accuracy", "tie at 0.7747", "near-identically"],
        "quotes": [
            "Among image encoders, DINOv3-L and EVA-02-L (Fang et al., 2024) tie at 0.7747, DINOv2-L "
            "(Oquab et al., 2024) is marginally lower at 0.7736, and the efficiency-oriented Hiera-L "
            "(Ryali et al., 2023) is the weakest in all four pairings.",
            "Because the top two image encoders are indistinguishable on accuracy, the selection of DINOv3-L "
            "for deployment rests on recency and licensing/maintenance considerations rather than on a "
            "measured accuracy advantage, and we report it as such.",
            "Rather than committing to one encoder pair by convention, this study runs a full 4×4 ablation "
            "matrix (16 combinations) and documents that DINOv3-Large and EVA-02-Large perform "
            "near-identically while Hiera-Large is consistently the weakest across every text-encoder "
            "pairing — a design-space mapping rarely reported this completely in the citizen-reporting "
            "literature.",
        ],
    },
    "C4": {
        "title": "The value of PGS lies in its uncertainty signal, not in accuracy",
        "sections": "Sections 3.3, 4, 5.1",
        "phrases": ["uncertainty signal at negligible", "lies in providing a low-cost uncertainty signal",
                    "true value of PGS"],
        "quotes": [
            "The documented value of PGS in this system is therefore an uncertainty signal at negligible "
            "inference cost (about 0.05 ms per sample for the head with 30 virtual members), not an accuracy "
            "gain; if stronger reliability guarantees are required, coverage-guaranteed selective prediction "
            "or set-valued outputs through conformal prediction (Angelopoulos & Bates, 2023) are more "
            "appropriate than enlarging the virtual ensemble (see also Abdar et al., 2021; Gawlikowski et "
            "al., 2023).",
            "The value of PGS in this setting lies in providing a low-cost uncertainty signal, not in "
            "improving accuracy, because the dominant uncertainty in citizen reports is aleatoric.",
            "The true value of PGS lies in its function as a signal of epistemic uncertainty, rather than as "
            "a source of improved accuracy.",
        ],
    },
    "C5": {
        "title": "The export reproduces the manuscript numbers (Table 4 and encoder ablation)",
        "sections": "Section 3.3, Table 4; Section 3.2, Figure 4",
        "phrases": ["Table 4. Separation"],
        "quotes": ["Table 4. Separation of the tree-count effect from the PGS effect (DINOv3-L + mE5-L, "
                   "test set, n = 9,266)."],
    },
    "C6": {
        "title": "Top-1 predictions of pgs/argmax and pgs/pgs differ on 130 test samples",
        "sections": "Sections 3.4, 4, 5.1; Table 5 note; Table 10",
        "phrases": ["130 of 9,266", "130 out of 9,266", "(130 samples)"],
        "quotes": [
            "Top-1 predictions of the native PGS path and the ONNX path differ on 130 of 9,266 samples.",
            "The system that is actually served therefore produces non-PGS probabilities: on 130 of 9,266 "
            "test samples its top-1 prediction differs from the PGS path that is reported as the headline "
            "result, and its accuracy and macro-F1 (0.8116 and 0.7793) are 0.43 and 0.46 pp higher.",
            "From an engineering perspective, the classifier head was exported to ONNX with numerical parity "
            "(|Δp|max = 1.38 × 10−6), but PGS does not survive the export, so the server model differs from "
            "the reported one on 130 of 9,266 test samples and lacks the uncertainty signal; documenting and "
            "resolving such train–serving gaps is necessary for trustworthy deployment.",
            "Consequently, the model used in production differs from the one reported as the primary result: "
            "130 out of 9,266 test samples yielded different top-1 predictions, with the production pipeline "
            "recording a higher score (accuracy of 0.8116 versus 0.8074).",
        ],
    },
}


@functools.lru_cache(maxsize=4)
def _docx_texts(docx_path: str) -> tuple[str, ...]:
    """Paragraph texts and de-duplicated table-row texts of a DOCX (read-only, cached)."""
    import docx  # python-docx; imported lazily because it is only needed for --manuscript

    doc = docx.Document(docx_path)
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            cells: list[str] = []
            for c in r.cells:
                if c.text not in cells:
                    cells.append(c.text)
            texts.append(" | ".join(cells))
    return tuple(texts)


#: Sentence boundary: ., ! or ? followed by whitespace and a capital/opening bracket, but not after
#: common abbreviations ("vs.", "et al.", "e.g.", "i.e.", "cf.", "Fig.").
_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])(?<!\bvs\.)(?<!\bal\.)(?<!e\.g\.)(?<!i\.e\.)(?<!\bcf\.)(?<!\bFig\.)"
                             r"\s+(?=[A-Z(\"“])")


def manuscript_sentences(docx_path: Path, phrases: Sequence[str]) -> list[str]:
    """Sentences of a DOCX (paragraphs and table rows) containing any phrase. Read-only."""
    found: list[str] = []
    for text in _docx_texts(str(docx_path)):
        for sent in _SENTENCE_SPLIT.split(text):
            if any(ph.lower() in sent.lower() for ph in phrases) and sent.strip() not in found:
                found.append(sent.strip())
    return found


def evaluate_claims(results: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Verdicts for C1-C6 following CONTRACT.md section 5 exactly."""
    pr = {r["row"]: r for r in results["paired"]["rows"]}
    supp = results["paired"]["supplementary"][0]
    sel = {r["score"]: r for r in results["selective"]["rows"]}
    rep = results["reproduction"]
    claims: list[dict[str, Any]] = []

    # C1
    r4 = pr["R4"]
    if r4["d_accuracy"] < 0 and r4["p_mcnemar"] < ALPHA:
        v1 = "SUPPORTED"
    elif r4["d_accuracy"] < 0:
        v1 = "QUALIFIED"
    else:
        v1 = "CONTRADICTED"
    ev1 = (f"R4 (pgs/pgs − pgs/argmax): Δaccuracy = {fmt_pp(r4['d_accuracy'])} pp "
           f"{fmt_pp_ci(*r4['d_accuracy_ci'])}, Δmacro-F1 = {fmt_pp(r4['d_macro_f1'])} pp "
           f"{fmt_pp_ci(*r4['d_macro_f1_ci'])}; b / c = {fmt_int(r4['b'])} / {fmt_int(r4['c'])}; McNemar "
           f"{p_phrase(r4['p_mcnemar'])} (Holm {fmt_p(r4['p_holm'])}). Supplementary, baseline checkpoint "
           f"(cb/pgs − cb/argmax): Δaccuracy = {fmt_pp(supp['d_accuracy'])} pp "
           f"{fmt_pp_ci(*supp['d_accuracy_ci'])}, Δmacro-F1 = {fmt_pp(supp['d_macro_f1'])} pp, McNemar "
           f"{p_phrase(supp['p_mcnemar'])}.")
    if v1 == "SUPPORTED":
        s1 = None
    elif v1 == "QUALIFIED":
        s1 = (f"At the fixed PGS checkpoint, PGS-averaged inference changes accuracy by "
              f"{fmt_pp(r4['d_accuracy'])} pp (95% CI {fmt_pp_ci(*r4['d_accuracy_ci'])}; McNemar "
              f"{p_phrase(r4['p_mcnemar'])}), a reduction that is not statistically significant.")
    else:
        verb1 = ("raises accuracy significantly" if r4["d_accuracy"] > 0 and r4["p_mcnemar"] < ALPHA
                 else "does not lower accuracy")
        s1 = (f"At the fixed PGS checkpoint, PGS-averaged inference {verb1} "
              f"(Δaccuracy = {fmt_pp(r4['d_accuracy'])} pp, 95% CI {fmt_pp_ci(*r4['d_accuracy_ci'])}; "
              f"McNemar {p_phrase(r4['p_mcnemar'])}).")
    claims.append({"id": "C1", "verdict": v1,
                   "criterion": "R4 Δaccuracy < 0 and McNemar p < 0.05 (QUALIFIED: negative but not "
                                "significant; CONTRADICTED: Δaccuracy ≥ 0).",
                   "evidence": ev1, "suggestion": s1, "refinement": None})

    # C2
    r5 = pr["R5"]
    mc_sig = r5["p_mcnemar"] < ALPHA
    ci_ex = ci_excludes_zero(*r5["d_macro_f1_ci"])
    v2 = "CONTRADICTED" if (mc_sig or ci_ex) else "SUPPORTED"
    ev2 = (f"R5 (pgs/pgs − cb/argmax): Δmacro-F1 = {fmt_pp(r5['d_macro_f1'])} pp "
           f"{fmt_pp_ci(*r5['d_macro_f1_ci'])} ({'excludes' if ci_ex else 'includes'} 0); Δaccuracy = "
           f"{fmt_pp(r5['d_accuracy'])} pp {fmt_pp_ci(*r5['d_accuracy_ci'])}; b / c = {fmt_int(r5['b'])} / "
           f"{fmt_int(r5['c'])}; McNemar {p_phrase(r5['p_mcnemar'])} (Holm {fmt_p(r5['p_holm'])}).")
    if v2 == "CONTRADICTED":
        # Attribution follows from the decomposition R5 = R4 (inference mode, same checkpoint) + checkpoint
        # change (pgs/argmax - cb/argmax, i.e. more boosting iterations).  A non-significant positive R4 is
        # not evidence that PGS averaging contributes nothing, so it is phrased as "cannot be attributed".
        r4_f1 = (f"(Δmacro-F1 = {fmt_pp(r4['d_macro_f1'])} pp, 95% CI {fmt_pp_ci(*r4['d_macro_f1_ci'])})")
        if r5["d_macro_f1"] <= 0:
            attribution = (f"at a fixed checkpoint, PGS-averaged inference changes macro-F1 by "
                           f"{fmt_pp(r4['d_macro_f1'])} pp (95% CI {fmt_pp_ci(*r4['d_macro_f1_ci'])})")
        elif r4["d_macro_f1"] <= 0:
            attribution = (f"because PGS-averaged inference does not raise macro-F1 at a fixed checkpoint {r4_f1}, "
                           "the difference arises from the change of checkpoint (trained with more boosting "
                           "iterations) rather than from PGS-averaged inference")
        elif not ci_excludes_zero(*r4["d_macro_f1_ci"]):
            attribution = (f"PGS-averaged inference does not raise macro-F1 significantly at a fixed checkpoint "
                           f"{r4_f1}, so the difference cannot be attributed to PGS-averaged inference")
        else:
            attribution = ("part of it is attributable to PGS averaging itself, which raises macro-F1 at a fixed "
                           f"checkpoint by {fmt_pp(r4['d_macro_f1'])} pp (95% CI "
                           f"{fmt_pp_ci(*r4['d_macro_f1_ci'])})")
        holm_sig = r5["p_holm"] < ALPHA
        holm_c2 = "" if holm_sig else f"; Holm-adjusted {p_phrase(r5['p_holm'])}"
        if mc_sig and ci_ex:
            s2 = (f"A paired analysis on identical test samples shows that the naive "
                  f"{fmt_pp(r5['d_macro_f1'])} pp macro-F1 difference is statistically significant "
                  f"(95% CI of Δmacro-F1 {fmt_pp_ci(*r5['d_macro_f1_ci'])}; McNemar {p_phrase(r5['p_mcnemar'])}"
                  f"{holm_c2}); {attribution}.")
        else:
            which = (f"McNemar's test on top-1 correctness is {'' if holm_sig else 'only nominally '}significant "
                     f"({p_phrase(r5['p_mcnemar'])}{holm_c2}) but the bootstrap interval of Δmacro-F1 "
                     f"{fmt_pp_ci(*r5['d_macro_f1_ci'])} includes zero") if mc_sig else (
                     f"the bootstrap interval of Δmacro-F1 {fmt_pp_ci(*r5['d_macro_f1_ci'])} excludes zero "
                     f"but McNemar's test on top-1 correctness is not significant ({p_phrase(r5['p_mcnemar'])})")
            s2 = (f"Whether the naive {fmt_pp(r5['d_macro_f1'])} pp macro-F1 difference exceeds sampling "
                  f"variance is not settled by the paired analysis: {which}; {attribution}.")
        ref2 = None
    else:
        s2 = None
        ref2 = ("Replace the unpaired argument (a point value inside the other configuration's interval) "
                f"with the paired result: \"A paired bootstrap on identical resamples gives Δmacro-F1 = "
                f"{fmt_pp(r5['d_macro_f1'])} pp (95% CI {fmt_pp_ci(*r5['d_macro_f1_ci'])}), and McNemar's exact "
                f"test gives {p_phrase(r5['p_mcnemar'])}, so the naive difference cannot be distinguished "
                "from sampling variance.\"")
    claims.append({"id": "C2", "verdict": v2,
                   "criterion": "CONTRADICTED if R5 McNemar p < 0.05 or the Δmacro-F1 95% CI excludes 0.",
                   "evidence": ev2, "suggestion": s2, "refinement": ref2})

    # C3
    r1 = pr["R1"]
    holm_ns = r1["p_holm"] >= ALPHA
    f1_inc = not ci_excludes_zero(*r1["d_macro_f1_ci"])
    if holm_ns and f1_inc:
        v3 = "SUPPORTED"
    elif not holm_ns and not f1_inc:
        v3 = "CONTRADICTED"
    else:
        v3 = "QUALIFIED"
    ev3 = (f"R1 ({r1['label_a']} − {r1['label_b']}): Δaccuracy = {fmt_pp(r1['d_accuracy'])} pp "
           f"{fmt_pp_ci(*r1['d_accuracy_ci'])}, Δmacro-F1 = {fmt_pp(r1['d_macro_f1'])} pp "
           f"{fmt_pp_ci(*r1['d_macro_f1_ci'])}; b / c = {fmt_int(r1['b'])} / {fmt_int(r1['c'])}; McNemar "
           f"{p_phrase(r1['p_mcnemar'])}, Holm-adjusted {p_phrase(r1['p_holm'])}. Caveat: a non-significant "
           "difference is not an equivalence test.")
    e_a, e_b = display(r1["key_a"][0].split("__")[0]), display(r1["key_b"][0].split("__")[0])
    if v3 == "SUPPORTED":
        s3 = None
        ref3 = (f"Qualify \"indistinguishable\" with the paired test and its caveat: \"On identical test samples, "
                f"the difference between {e_a} and {e_b} is not statistically significant (McNemar's exact test, "
                f"Holm-adjusted {p_phrase(r1['p_holm'])}; Δmacro-F1 = {fmt_pp(r1['d_macro_f1'])} pp, 95% CI "
                f"{fmt_pp_ci(*r1['d_macro_f1_ci'])}); this is not a formal equivalence test, so it does not "
                "show that the two encoders perform identically.\"")
    elif v3 == "CONTRADICTED":
        s3 = (f"{e_a} and {e_b} differ significantly on the test set (Δmacro-F1 = {fmt_pp(r1['d_macro_f1'])} pp, "
              f"95% CI {fmt_pp_ci(*r1['d_macro_f1_ci'])}; McNemar's exact test, Holm-adjusted "
              f"{p_phrase(r1['p_holm'])}).")
        ref3 = None
    else:
        part = ("McNemar's test on top-1 correctness is significant after Holm adjustment "
                f"({p_phrase(r1['p_holm'])}) while the Δmacro-F1 interval {fmt_pp_ci(*r1['d_macro_f1_ci'])} "
                "includes zero") if not holm_ns else (
                f"the Δmacro-F1 interval {fmt_pp_ci(*r1['d_macro_f1_ci'])} excludes zero while McNemar's test "
                f"is not significant after Holm adjustment ({p_phrase(r1['p_holm'])})")
        s3 = (f"The paired analysis gives a mixed picture for {e_a} versus {e_b}: {part}; the two encoders "
              "should therefore be described as close rather than indistinguishable.")
        ref3 = None
    claims.append({"id": "C3", "verdict": v3,
                   "criterion": "SUPPORTED if R1 Holm-adjusted p ≥ 0.05 and the Δmacro-F1 CI includes 0 "
                                "(CONTRADICTED if neither holds; QUALIFIED if exactly one holds).",
                   "evidence": ev3, "suggestion": s3, "refinement": ref3})

    # C4
    s5, s1 = sel["S5"], sel["S1"]
    above = s5["auroc_ci"][0] > 0.5
    includes_half = s5["auroc_ci"][0] <= 0.5 <= s5["auroc_ci"][1]
    overlap = intervals_overlap(s5["auroc_ci"], s1["auroc_ci"])
    clearly_worse = (not overlap) and s5["auroc"] < s1["auroc"]
    if includes_half or not above:
        v4 = "CONTRADICTED"
    elif clearly_worse:
        v4 = "QUALIFIED"
    else:
        v4 = "SUPPORTED"
    d51 = next(d for d in results["selective"]["paired_auroc_differences"] if d["a"] == "S5" and d["b"] == "S1")
    ev4 = (f"S5 (mutual information) AUROC = {fmt_metric_ci(s5['auroc'], *s5['auroc_ci'])}; S1 "
           f"(1 − max p, served argmax) AUROC = {fmt_metric_ci(s1['auroc'], *s1['auroc_ci'])}; intervals "
           f"{'overlap' if overlap else 'do not overlap'}. Supplementary paired bootstrap ΔAUROC (S5 − S1) = "
           f"{fmt_fixed(d51['d_auroc'], 4, plus=True)} {fmt_ci(*d51['d_auroc_ci'])}. R4 Δaccuracy = "
           f"{fmt_pp(r4['d_accuracy'])} pp (McNemar {p_phrase(r4['p_mcnemar'])}).")
    if v4 == "SUPPORTED":
        s4 = None
        r4_stat = f"Δaccuracy = {fmt_pp(r4['d_accuracy'])} pp, McNemar {p_phrase(r4['p_mcnemar'])}"
        if r4["p_mcnemar"] >= ALPHA:
            acc4 = f"At a fixed checkpoint PGS averaging does not change accuracy significantly ({r4_stat}), whereas"
        elif r4["d_accuracy"] < 0:
            acc4 = f"At a fixed checkpoint PGS averaging lowers accuracy ({r4_stat}), whereas"
        else:
            acc4 = f"At a fixed checkpoint PGS averaging raises accuracy ({r4_stat}), and in addition"
        ref4 = (f"Back the claim with the per-sample evidence: \"{acc4} its mutual information separates correct "
                f"from incorrect routings (AUROC = {fmt_metric(s5['auroc'])}, 95% CI {fmt_ci(*s5['auroc_ci'])}).\"")
        if ci_excludes_zero(*d51["d_auroc_ci"]) and d51["d_auroc"] < 0:
            ref4 += (" Do not present mutual information as a better error signal than the served path: the "
                     "paired bootstrap difference to 1 − max p of the served argmax path excludes zero "
                     f"(ΔAUROC = {fmt_fixed(d51['d_auroc'], 4, plus=True)}, 95% CI {fmt_ci(*d51['d_auroc_ci'])}), "
                     "although the marginal intervals overlap.")
        if r4["p_mcnemar"] < ALPHA and r4["d_accuracy"] > 0:
            ref4 += (" Drop \"not in accuracy\": PGS averaging significantly improves accuracy at a fixed "
                     "checkpoint.")
    elif v4 == "QUALIFIED":
        s4 = ("The mutual information from PGS carries an error signal (AUROC = "
              f"{fmt_metric(s5['auroc'])}, 95% CI {fmt_ci(*s5['auroc_ci'])}), but the maximum class probability "
              "of the served argmax path, which requires no PGS, detects misrouted reports better (AUROC = "
              f"{fmt_metric(s1['auroc'])}, 95% CI {fmt_ci(*s1['auroc_ci'])}); PGS is therefore not required "
              "for uncertainty-based deferral in this system.")
        ref4 = None
    else:
        s4 = ("The mutual information from PGS does not separate correct from incorrect predictions better "
              f"than chance (AUROC = {fmt_metric(s5['auroc'])}, 95% CI {fmt_ci(*s5['auroc_ci'])}); the claimed "
              "value of PGS as an uncertainty signal is not supported by the per-sample analysis.")
        ref4 = None
    claims.append({"id": "C4", "verdict": v4,
                   "criterion": "SUPPORTED if the S5 AUROC CI lies above 0.5 and S5 is not clearly worse than "
                                "S1; QUALIFIED if above 0.5 but lower than S1 with non-overlapping CIs; "
                                "CONTRADICTED if the S5 CI includes 0.5.",
                   "evidence": ev4, "suggestion": s4, "refinement": ref4})

    # C5
    flagged = [r for r in rep["rows"] if r["reproduces_manifest"] is not None or r["reproduces_computed"] is not None]
    fails = [r for r in flagged if r["reproduces_manifest"] is False or r["reproduces_computed"] is False]
    inconsistent = [r for r in flagged if r["reproduces_manifest"] is not None and r["reproduces_computed"] is not None
                    and r["reproduces_manifest"] != r["reproduces_computed"]]
    if fails:
        v5 = "CONTRADICTED"
    elif inconsistent or not flagged:
        v5 = "QUALIFIED"
    else:
        v5 = "SUPPORTED"
    parts = []
    for r in flagged:
        chk = "; ".join(f"{m} {fmt_metric(c['value'])} (expected {' or '.join(fmt_metric(x) for x in c['expected'])})"
                        for m, c in r["checks"].items())
        parts.append(f"{r['pair']} {r['checkpoint']}/{r['mode']}: {chk}; manifest flag = "
                     f"{r['reproduces_manifest']}, recomputed flag = {r['reproduces_computed']}")
    ev5 = "\n".join(f"- {x}" for x in parts) if parts else "No manuscript expectations were found in the manifest."
    if fails:
        s5_ = ("Update the affected Table 4 / Figure 4 values to the exported ones or re-export from the exact "
               "checkpoints used for the manuscript: " + "; ".join(
                   f"{r['pair']} {r['checkpoint']}/{r['mode']}: accuracy {fmt_metric(r['accuracy'])}, "
                   f"macro-F1 {fmt_metric(r['macro_f1'])}" for r in fails) + ".")
    elif not flagged:
        # e.g. an export run with --no-expect: nothing to compare, so say how to get the check.
        s5_ = ("Re-run export_per_sample_predictions.py without --no-expect on the RunPod artifacts so that the "
               "manifest records the manuscript values, then re-run this analysis; do not rely on TXT_REPRO "
               "until C5 is SUPPORTED.")
    elif inconsistent:
        s5_ = ("The manifest flags and the values recomputed from the CSVs disagree; re-export before using "
               "these results.")
    else:
        s5_ = None
    claims.append({"id": "C5", "verdict": v5, "suggestion_is_instruction": True,
                   "criterion": "Manifest reproduces_manuscript is true for every row with an expectation "
                                "(cross-checked by recomputing accuracy and macro-F1 from the CSVs).",
                   "evidence": ev5, "suggestion": s5_, "refinement": None})

    # C6
    n_dis = rep["top1_disagreements_pgs_vs_argmax"]
    v6 = "SUPPORTED" if n_dis == 130 else "CONTRADICTED"
    ev6 = (f"Recomputed top-1 disagreements between pgs/argmax and pgs/pgs on the test set: {fmt_int(n_dis)} "
           f"(manifest: {rep['top1_disagreements_manifest']}; manuscript: 130).")
    s6 = None if v6 == "SUPPORTED" else (f"Replace \"130\" by \"{fmt_int(n_dis)}\" in every sentence listed "
                                         "above (and in the Table 5 note).")
    claims.append({"id": "C6", "verdict": v6, "criterion": "Top-1 disagreements = 130.", "suggestion_is_instruction": True,
                   "evidence": ev6, "suggestion": s6, "refinement": None})

    for c in claims:
        src = CLAIM_SOURCES[c["id"]]
        c.update({"title": src["title"], "sections": src["sections"]})
    return claims


# ---------------------------------------------------------------------------------------------
# Fill values and generated prose (CONTRACT.md section 4)
# ---------------------------------------------------------------------------------------------


def _conf_state(conf: float, acc: float) -> tuple[str, str]:
    """('over'|'under'|'matched', gap in pp formatted without sign) using the 4-dp table values."""
    gap = round(conf, 4) - round(acc, 4)
    gap_s = fmt_fixed(abs(gap) * 100.0, 2)
    if gap_s == "0.00":
        return "matched", gap_s
    return ("over" if gap > 0 else "under"), gap_s


def sig_phrase(p: float, p_holm: float) -> str:
    """Significance statement with raw and Holm-adjusted McNemar p."""
    raw_sig, holm_sig = p < ALPHA, p_holm < ALPHA
    if raw_sig and holm_sig:
        return f"statistically significant (McNemar {p_phrase(p)}; Holm-adjusted {p_phrase(p_holm)})"
    if raw_sig:
        return (f"nominally significant (McNemar {p_phrase(p)}) but not significant after Holm adjustment "
                f"({p_phrase(p_holm)})")
    return f"not statistically significant (McNemar {p_phrase(p)}; Holm-adjusted {p_phrase(p_holm)})"


def txt_repro(results: Mapping[str, Any], deployed: str) -> str:
    """Fragment continuing "The exported records reproduce the accuracy and macro-F1 values of Table 4 "."""
    rows = [r for r in results["reproduction"]["rows"] if r["pair"] == deployed]
    order = [("cb", "argmax"), ("cb", "pgs"), ("pgs", "argmax"), ("pgs", "pgs")]
    rows = sorted(rows, key=lambda r: order.index((r["checkpoint"], r["mode"])))
    with_exp = [r for r in rows if r["checks"]]
    no_exp = [r for r in rows if not r["checks"]]
    ok = [r for r in with_exp if all(c["matches"] for c in r["checks"].values())]
    bad = [r for r in with_exp if r not in ok]
    n_total = len(rows)

    def deviation(r: Mapping[str, Any]) -> str:
        bad_bits, ok_bits = [], []
        for metric, label, name in (("accuracy", "an accuracy", "accuracy"), ("macro_f1", "a macro-F1", "macro-F1")):
            c = r["checks"].get(metric)
            if c is None:
                continue
            exp_s = " or ".join(fmt_metric(x) for x in c["expected"])
            if c["matches"]:
                ok_bits.append(f"the {name} ({fmt_metric(c['value'])}) matches")
            else:
                bad_bits.append(f"{label} of {fmt_metric(c['value'])} instead of {exp_s}")
        s = f"{join_list(bad_bits)} for {COMBO_PROSE[(r['checkpoint'], r['mode'])]}"
        return s + (f", whereas {ok_bits[0]}" if ok_bits else "")

    max_dev = max((abs(c["deviation_pp"]) for r in bad for c in r["checks"].values() if not c["matches"]),
                  default=0.0)
    dev_note = f" (maximum deviation {fmt_fixed(max_dev, 2)} pp)" if bad else ""
    missing = (f"; no manuscript value was available for {join_list([COMBO_PROSE[(r['checkpoint'], r['mode'])] for r in no_exp])}"
               if no_exp else "")
    # Every branch must read as a grammatical continuation of the fixed sentence opening
    # "The exported records reproduce the accuracy and macro-F1 values of Table 4 ".
    if n_total == 0:
        return "to an extent that cannot be checked here, because no records of the deployed pair were exported."
    if not bad and len(ok) == n_total and n_total == 4:
        return "to the fourth decimal place for all four checkpoint–mode combinations."
    if not bad and ok:
        return (f"to the fourth decimal place for all {number_word(len(ok))} checkpoint–mode combinations "
                f"with a reported value{missing}.")
    if not with_exp:
        return ("to an extent that cannot be verified here, because the manifest lists no manuscript values "
                "for comparison.")
    if ok:
        return (f"to the fourth decimal place for {number_word(len(ok))} of the {number_word(len(with_exp))} "
                f"checkpoint–mode combinations but give {'; '.join(deviation(r) for r in bad)}"
                f"{dev_note}{missing}.")
    return (f"only approximately: they give {'; '.join(deviation(r) for r in bad)}"
            f"{dev_note}{missing}.")


def _members(results: Mapping[str, Any]) -> int:
    """Number of virtual ensemble members M (manifest ``virtual_ensembles``; contract default 30)."""
    m = ((results.get("reproduction") or {}).get("manifest_meta") or {}).get("virtual_ensembles")
    return int(m) if _finite(m) else 30


def txt_calibration(results: Mapping[str, Any]) -> str:
    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    r3, r4, r5 = cal["R3"], cal["R4"], cal["R5"]
    temp = results["calibration"]["temperature"]
    t = temp["temperature"]
    m = _members(results)
    e3, e4, e5 = fmt_metric(r3["ece"]), fmt_metric(r4["ece"]), fmt_metric(r5["ece"])
    c3, c4, c5 = fmt_ci(*r3["ece_ci"]), fmt_ci(*r4["ece_ci"]), fmt_ci(*r5["ece_ci"])
    sents = []
    if intervals_overlap(r3["ece_ci"], r4["ece_ci"]):
        # Overlapping intervals are absence of evidence, not evidence of unchanged calibration.
        sents.append(f"At the fixed PGS checkpoint, PGS-averaged inference (M = {m}) does not change calibration "
                     f"detectably relative to the served argmax path: the ECE is {e4} (95% CI {c4}) "
                     f"versus {e3} ({c3}), and the two bootstrap intervals overlap.")
    elif r4["ece"] < r3["ece"]:
        sents.append(f"At the fixed PGS checkpoint, PGS-averaged inference (M = {m}) improves calibration "
                     f"relative to the served argmax path, lowering the ECE from {e3} (95% CI {c3}) to {e4} "
                     f"({c4}) with non-overlapping bootstrap intervals.")
    else:
        sents.append(f"At the fixed PGS checkpoint, PGS-averaged inference (M = {m}) worsens calibration "
                     f"relative to the served argmax path, raising the ECE from {e3} (95% CI {c3}) to {e4} "
                     f"({c4}) with non-overlapping bootstrap intervals.")
    db = cmp_formatted(r4["brier"], r3["brier"], fmt_metric)
    dn = cmp_formatted(r4["nll"], r3["nll"], fmt_metric)
    b3, b4, n3, n4 = (fmt_metric(r3["brier"]), fmt_metric(r4["brier"]), fmt_metric(r3["nll"]), fmt_metric(r4["nll"]))
    if db == dn == 0:
        sents.append(f"The Brier score ({b3}) and the NLL ({n3} nats) are identical to four decimal places "
                     "under both inference modes.")
    elif db == dn:
        verb = "lowers" if db < 0 else "raises"
        marginal = (abs(r4["brier"] - r3["brier"]) < 0.01 * r3["brier"]
                    and abs(r4["nll"] - r3["nll"]) < 0.01 * r3["nll"])
        adverb = " only marginally" if marginal else ""
        sents.append(f"PGS averaging {verb} both the Brier score (from {b3} to {b4}) and the NLL "
                     f"(from {n3} to {n4} nats){adverb}.")
    else:
        def change(d: int, a: str, b: str) -> str:
            if d == 0:
                return f"remains at {a}"
            return f"{'decreases' if d < 0 else 'increases'} from {a} to {b}"
        sents.append(f"Under PGS averaging, the Brier score {change(db, b3, b4)}, whereas the NLL "
                     f"{change(dn, n3, n4)} nats.")
    s3_, g3 = _conf_state(r3["mean_confidence"], r3["accuracy"])
    s4_, g4 = _conf_state(r4["mean_confidence"], r4["accuracy"])
    if s3_ == s4_ == "over":
        sents.append(f"Both inference paths are overconfident on average, with mean confidence exceeding "
                     f"accuracy by {g3} pp (served argmax) and {g4} pp (PGS-averaged).")
    elif s3_ == s4_ == "under":
        sents.append(f"Both inference paths are underconfident on average, with mean confidence falling short "
                     f"of accuracy by {g3} pp (served argmax) and {g4} pp (PGS-averaged).")
    elif s3_ == s4_ == "matched":
        sents.append("On average, mean confidence matches accuracy to within 0.01 pp for both inference paths.")
    else:
        def state(s: str, g: str) -> str:
            return {"over": f"overconfident on average (mean confidence exceeds accuracy by {g} pp)",
                    "under": f"underconfident on average (mean confidence falls short of accuracy by {g} pp)",
                    "matched": "neither over- nor underconfident on average"}[s]
        sents.append(f"The served argmax path is {state(s3_, g3)}, whereas PGS-averaged inference is "
                     f"{state(s4_, g4)}.")
    ts = fmt_temp(t)
    d_ts = cmp_formatted(r5["ece"], r3["ece"], fmt_metric)
    ov = intervals_overlap(r5["ece_ci"], r3["ece_ci"])
    if d_ts < 0 and not ov:
        eff = f"reduces the ECE to {e5} (95% CI {c5}), an interval that does not overlap that of the unscaled path"
    elif d_ts < 0:
        eff = f"reduces the ECE to {e5} (95% CI {c5}), although the two intervals overlap"
    elif d_ts > 0:
        eff = f"does not reduce the ECE on the test set ({e5}, 95% CI {c5})"
    else:
        eff = f"keeps the ECE at {e5} (95% CI {c5})"
    if ts == "1.00":
        sents.append(f"Temperature scaling fitted on the validation set (T = {ts}) leaves the served "
                     f"probabilities practically unchanged and {eff}.")
    else:
        shape = "softens" if t > 1 else "sharpens"
        # A positive temperature preserves the ranking of the classes; only written-precision ties can flip.
        keep = " without changing any top-1 prediction" if not temp.get("test_argmax_changes") else ""
        sents.append(f"Temperature scaling fitted on the validation set (T = {ts}) {shape} the served "
                     f"probabilities{keep} and {eff}.")
    return " ".join(sents)


def _encoder_names(r: Mapping[str, Any]) -> tuple[str, str]:
    return display(r["key_a"][0].split("__")[0]), display(r["key_b"][0].split("__")[0])


def _encoder_stat(r: Mapping[str, Any]) -> str:
    return f"Δmacro-F1 = {fmt_pp(r['d_macro_f1'])} pp, adjusted {p_phrase(r['p_holm'])}"


def _encoder_item(r: Mapping[str, Any]) -> str:
    a, b = _encoder_names(r)
    return f"{a} vs. {b}: {_encoder_stat(r)}"


def txt_paired(results: Mapping[str, Any]) -> str:
    pr = {r["row"]: r for r in results["paired"]["rows"]}
    enc = [pr["R1"], pr["R2"], pr["R3"]]
    txt_enc = display(enc[0]["key_a"][0].split("__", 1)[1])
    sig = [r for r in enc if r["p_holm"] < ALPHA]
    ns = [r for r in enc if r["p_holm"] >= ALPHA]
    sents = []
    lead = f"Among the image encoders, each fused with {txt_enc} and evaluated with PGS-averaged inference"
    if not sig:
        sents.append(f"{lead}, none of the three pairwise differences is significant after Holm adjustment "
                     f"(McNemar's exact test; {'; '.join(_encoder_item(r) for r in enc)}), although a "
                     "non-significant difference is not evidence of equivalence.")
    elif not ns:
        sents.append(f"{lead}, all three pairwise differences are significant after Holm adjustment "
                     f"(McNemar's exact test; {'; '.join(_encoder_item(r) for r in enc)}).")
    else:
        sep = join_list([f"{a} from {b} ({_encoder_stat(r)})" for r in sig for a, b in [_encoder_names(r)]])
        rest = join_list([f"between {a} and {b} ({_encoder_stat(r)})" for r in ns for a, b in [_encoder_names(r)]])
        noun = "difference" if len(ns) == 1 else "differences"
        verb = "is" if len(ns) == 1 else "are"
        sents.append(f"{lead}, McNemar's exact test with Holm adjustment separates {sep}, whereas the {noun} "
                     f"{rest} {verb} not significant; a non-significant difference is, however, not evidence of "
                     "equivalence.")
    discord = [r for r in enc if (r["p_holm"] < ALPHA) != ci_excludes_zero(*r["d_macro_f1_ci"])]
    if discord:
        bits = []
        for r in discord:
            a, b = _encoder_names(r)
            if r["p_holm"] < ALPHA:
                bits.append(f"for {a} vs. {b}, the Δmacro-F1 interval {fmt_pp_ci(*r['d_macro_f1_ci'])} "
                            "includes zero although McNemar's test is significant")
            else:
                bits.append(f"for {a} vs. {b}, the Δmacro-F1 interval {fmt_pp_ci(*r['d_macro_f1_ci'])} "
                            "excludes zero although McNemar's test is not significant")
        sents.append("The two paired analyses do not fully agree: " + "; ".join(bits) + ".")
    r4 = pr["R4"]
    sents.append(f"At the fixed PGS checkpoint, PGS-averaged inference changes accuracy by "
                 f"{fmt_pp(r4['d_accuracy'])} pp (95% CI {fmt_pp_ci(*r4['d_accuracy_ci'])}) and macro-F1 by "
                 f"{fmt_pp(r4['d_macro_f1'])} pp ({fmt_pp_ci(*r4['d_macro_f1_ci'])}) relative to the served argmax "
                 f"path; with {fmt_int(r4['b'])} reports correct only under PGS averaging and {fmt_int(r4['c'])} "
                 f"only under argmax, the accuracy difference is {sig_phrase(r4['p_mcnemar'], r4['p_holm'])}.")
    r5 = pr["R5"]
    mc_sig = r5["p_mcnemar"] < ALPHA
    ci_ex = ci_excludes_zero(*r5["d_macro_f1_ci"])
    df1 = fmt_pp(r5["d_macro_f1"])
    stat = (f"Δmacro-F1 = {df1} pp, 95% CI {fmt_pp_ci(*r5['d_macro_f1_ci'])}; McNemar {p_phrase(r5['p_mcnemar'])}, "
            f"Holm-adjusted {p_phrase(r5['p_holm'])}")
    what = (f"the naive {df1} pp macro-F1 difference between the PGS configuration and the baseline checkpoint "
            "with argmax inference")
    if mc_sig and ci_ex:
        holm_note = "" if r5["p_holm"] < ALPHA else ", although not after Holm adjustment"
        sents.append(f"Under the paired tests, {what} is statistically significant{holm_note} ({stat}).")
    elif not mc_sig and not ci_ex:
        sents.append(f"The paired tests confirm that {what} is not statistically significant ({stat}).")
    elif mc_sig and r5["p_holm"] >= ALPHA:
        # Only the unadjusted McNemar p is below alpha: calling it "significant" would overclaim.
        sents.append(f"Under the paired tests, {what} is not statistically significant after Holm adjustment: "
                     "McNemar's test on top-1 correctness is only nominally significant, and the bootstrap "
                     f"interval of Δmacro-F1 includes zero ({stat}).")
    elif mc_sig:
        sents.append(f"The paired tests disagree on {what}: McNemar's test on top-1 correctness is significant, "
                     f"but the bootstrap interval of Δmacro-F1 includes zero ({stat}).")
    else:
        sents.append(f"The paired tests disagree on {what}: the bootstrap interval of Δmacro-F1 excludes zero, "
                     f"but McNemar's test on top-1 correctness is not significant ({stat}).")
    return " ".join(sents)


def _best_score(rows: Sequence[Mapping[str, Any]]) -> tuple[Mapping[str, Any], Mapping[str, Any] | None]:
    finite = [r for r in rows if _finite(r["auroc"])]
    ranked = sorted(finite, key=lambda r: (-r["auroc"], r["row"]))
    return ranked[0], (ranked[1] if len(ranked) > 1 else None)


def _full_accuracy_str(results: Mapping[str, Any], score_row: Mapping[str, Any]) -> str:
    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    ref = cal["R3"] if score_row["score"] == "S1" else cal["R4"]
    return fmt_metric(ref["accuracy"])


def _deferral_clause(results: Mapping[str, Any], r: Mapping[str, Any]) -> tuple[int, str, str, str]:
    full_s = _full_accuracy_str(results, r)
    acc80 = fmt_metric(r["acc_at_80"])
    full_v = float(full_s)
    d = 0 if acc80 == full_s else (1 if r["acc_at_80"] > full_v else -1)
    # coverage_at_90 is the LARGEST k/N with accuracy >= 0.90; the curve need not be monotone, so the
    # wording must not claim that 90% holds at every smaller coverage.
    if r["k_at_90"] == 0:
        cov = "no coverage level reaches an accuracy of 90%"
    elif r["coverage_at_90"] >= 1.0:
        cov = "the accuracy is at least 90% even at full coverage"
    else:
        cov = (f"the largest coverage at which the accepted reports still reach an accuracy of at least 90% is "
               f"{fmt_cov(r['coverage_at_90'])}")
    return d, full_s, acc80, cov


def _paired_auroc(results: Mapping[str, Any], a: str, b: str) -> Mapping[str, Any] | None:
    return next((d for d in results["selective"].get("paired_auroc_differences", [])
                 if d["a"] == a and d["b"] == b), None)


def _mi_vs_served(results: Mapping[str, Any]) -> dict[str, Any]:
    """How S5 (MI) compares with S1 (1 - max p, served): marginal-CI overlap plus the paired ΔAUROC.

    Overlapping marginal intervals do not show that two AUROCs are equal; the paired bootstrap
    difference (same resamples, supplementary) is used to avoid calling them "comparable" when the
    paired interval excludes zero.  Returns direction ``sign`` (-1, 0, +1; 0 = no detectable
    difference) and a parenthetical with the paired statistic ("" when unavailable).
    """
    sel = {r["score"]: r for r in results["selective"]["rows"]}
    s5, s1 = sel["S5"], sel["S1"]
    overlap = intervals_overlap(s5["auroc_ci"], s1["auroc_ci"])
    d = _paired_auroc(results, "S5", "S1")
    paired = ""
    paired_excl = False
    if d is not None and _finite(d["d_auroc"]):
        paired = (f"paired bootstrap ΔAUROC = {fmt_fixed(d['d_auroc'], 4, plus=True)}, 95% CI "
                  f"{fmt_ci(*d['d_auroc_ci'])}")
        paired_excl = ci_excludes_zero(*d["d_auroc_ci"])
    if not overlap:
        sign = -1 if s5["auroc"] < s1["auroc"] else 1
    elif paired_excl:
        sign = -1 if d["d_auroc"] < 0 else 1
    else:
        sign = 0
    return {"overlap": overlap, "paired": paired, "paired_excludes_zero": paired_excl, "sign": sign}


def txt_selective(results: Mapping[str, Any]) -> str:
    rows = results["selective"]["rows"]
    sel = {r["score"]: r for r in rows}
    best, runner = _best_score(rows)
    sents = []
    head = (f"Among the six uncertainty scores, {SCORE_PROSE[best['score']]} has the highest error-detection "
            f"AUROC ({fmt_metric(best['auroc'])}, 95% CI {fmt_ci(*best['auroc_ci'])})")
    if runner is None:
        sents.append(head + ".")
    elif intervals_overlap(best["auroc_ci"], runner["auroc_ci"]):
        sents.append(head + f", although its interval overlaps that of {SCORE_PROSE[runner['score']]} "
                            f"({fmt_metric(runner['auroc'])}, {fmt_ci(*runner['auroc_ci'])}).")
    else:
        sents.append(head + f", ahead of {SCORE_PROSE[runner['score']]} ({fmt_metric(runner['auroc'])}, "
                            f"{fmt_ci(*runner['auroc_ci'])}) with non-overlapping intervals.")
    s5, s1 = sel["S5"], sel["S1"]
    a5, a1 = fmt_metric(s5["auroc"]), fmt_metric(s1["auroc"])
    c5, c1 = fmt_ci(*s5["auroc_ci"]), fmt_ci(*s1["auroc_ci"])
    mv = _mi_vs_served(results)
    lead5 = f"Mutual information reaches an AUROC of {a5} (95% CI {c5})"
    ref1 = f"that of {SCORE_PROSE['S1']} ({a1}, {c1})"
    if not mv["overlap"]:
        cmp_ = f"{lead5}, {'lower' if mv['sign'] < 0 else 'higher'} than {ref1} with non-overlapping intervals"
    elif mv["paired_excludes_zero"]:
        cmp_ = (f"{lead5}, slightly {'lower' if mv['sign'] < 0 else 'higher'} than {ref1}: the two intervals "
                f"overlap, but the paired difference excludes zero ({mv['paired']})")
    elif mv["paired"]:
        cmp_ = (f"{lead5}, not significantly different from {ref1}: the two intervals overlap and the paired "
                f"difference includes zero ({mv['paired']})")
    else:
        cmp_ = f"{lead5}; its interval overlaps {ref1}"
    lo5, hi5 = s5["auroc_ci"]
    if lo5 <= 0.5 <= hi5:
        cmp_ += ("; the interval for mutual information includes 0.5, so it does not separate correct from "
                 "incorrect predictions better than chance")
    elif hi5 < 0.5:
        cmp_ += ("; the interval for mutual information lies below 0.5, so higher mutual information goes with "
                 "fewer errors")
    sents.append(cmp_ + ".")
    d, full_s, acc80, cov = _deferral_clause(results, best)
    verb = {1: "raises", -1: "lowers", 0: "leaves"}[d]
    if d == 0:
        acc_part = f"leaves the accuracy of the accepted 80% unchanged at {acc80}"
    else:
        acc_part = f"{verb} the accuracy of the accepted 80% to {acc80} (from {full_s} at full coverage)"
    sents.append(f"Deferring the 20% least certain reports according to {SCORE_PROSE[best['score']]} {acc_part}, "
                 f"and {cov}.")
    return " ".join(sents)


def txt_implication(results: Mapping[str, Any], claims: Sequence[Mapping[str, Any]]) -> str:
    verdict = {c["id"]: c["verdict"] for c in claims}
    pr = {r["row"]: r for r in results["paired"]["rows"]}
    sel = {r["score"]: r for r in results["selective"]["rows"]}
    r4, s5, s1 = pr["R4"], sel["S5"], sel["S1"]
    dacc = fmt_pp(r4["d_accuracy"])
    if r4["d_accuracy"] > 0 and r4["p_mcnemar"] < ALPHA:
        acc_part = (f"contrary to the second half of the claim, PGS averaging improves accuracy at a fixed "
                    f"checkpoint (Δaccuracy = {dacc} pp)")
    elif r4["p_mcnemar"] < ALPHA:
        acc_part = f"PGS averaging does not improve accuracy at a fixed checkpoint but lowers it (Δaccuracy = {dacc} pp)"
    else:
        acc_part = (f"PGS averaging does not change accuracy significantly at a fixed checkpoint "
                    f"(Δaccuracy = {dacc} pp)")
    a5, c5, a1, c1 = fmt_metric(s5["auroc"]), fmt_ci(*s5["auroc_ci"]), fmt_metric(s1["auroc"]), fmt_ci(*s1["auroc_ci"])
    v = verdict["C4"]
    acc_improves = r4["d_accuracy"] > 0 and r4["p_mcnemar"] < ALPHA
    # Relation of MI to the PGS-free served score; never "comparable" on the strength of overlap alone.
    mv = _mi_vs_served(results)
    s1_ref = f"{SCORE_PROSE['S1']} ({a1}" + (f"; {mv['paired']})" if mv["paired"] else ")")
    if mv["sign"] > 0:
        rel = f" and {'' if not mv['overlap'] else 'slightly '}better than {s1_ref}"
    elif mv["sign"] < 0:
        rel = f", although slightly less well than {s1_ref}"
    elif mv["paired"]:
        rel = f", with no significant difference from {s1_ref}"
    else:
        rel = f", with an interval that overlaps that of {s1_ref}"
    if v == "SUPPORTED" and acc_improves:
        s_first = (f"Taken together, these results qualify the claim that the value of PGS lies in its uncertainty "
                   f"signal rather than in accuracy: PGS averaging does improve accuracy at a fixed checkpoint "
                   f"(Δaccuracy = {dacc} pp), and its mutual information separates correct from incorrect "
                   f"routings clearly better than chance (AUROC = {a5}, 95% CI {c5}){rel}.")
    elif v == "SUPPORTED":
        s_first = (f"Taken together, these results support the claim that the value of PGS lies in its uncertainty "
                   f"signal rather than in accuracy: {acc_part}, whereas its mutual information separates correct "
                   f"from incorrect routings clearly better than chance (AUROC = {a5}, 95% CI {c5}){rel}.")
    elif v == "QUALIFIED":
        s_first = (f"Taken together, these results qualify the claim that the value of PGS lies in its uncertainty "
                   f"signal rather than in accuracy: {acc_part}, and its mutual information does carry an error "
                   f"signal (AUROC = {a5}, 95% CI {c5}), but {SCORE_PROSE['S1']}, which requires no PGS, detects "
                   f"errors better ({a1}, {c1}).")
    else:
        if s5["auroc_ci"][1] < 0.5:
            chance = (f"the AUROC of its mutual information ({a5}, 95% CI {c5}) lies below chance, i.e. higher "
                      "mutual information goes with fewer errors")
        else:
            chance = f"the AUROC of its mutual information ({a5}, 95% CI {c5}) is not distinguishable from chance"
        s_first = (f"Taken together, these results do not support the claim that the value of PGS lies in its "
                   f"uncertainty signal: {acc_part}, and {chance}.")
    sents = [s_first]
    best, _ = _best_score(results["selective"]["rows"])
    d, full_s, acc80, _cov = _deferral_clause(results, best)
    if best["score"] == "S1":
        avail = ", which is already available in the served ONNX path"
    elif intervals_overlap(best["auroc_ci"], s1["auroc_ci"]):
        avail = (" (which requires native PGS inference that the ONNX export currently drops; "
                 f"{SCORE_PROSE['S1']}, whose AUROC interval overlaps, is a practical substitute that needs no PGS)")
    else:
        avail = ", which requires native PGS inference that the ONNX export currently drops"
    if d > 0:
        sents.append(f"Operationally, the 20% least certain reports could be deferred to a human operator on the "
                     f"basis of {SCORE_PROSE[best['score']]}{avail}; this would raise the accuracy of the "
                     f"automatically routed 80% from {full_s} to {acc80}.")
    else:
        sents.append(f"Operationally, deferring the 20% least certain reports according to "
                     f"{SCORE_PROSE[best['score']]} would not raise the accuracy of the automatically routed 80% "
                     f"({acc80} versus {full_s} at full coverage), so uncertainty-based deferral is not "
                     "recommended on this evidence.")
    return " ".join(sents)


def txt_conclusion(results: Mapping[str, Any], claims: Sequence[Mapping[str, Any]]) -> str:
    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    sel = {r["score"]: r for r in results["selective"]["rows"]}
    r3, r4, r5 = cal["R3"], cal["R4"], cal["R5"]
    e3, e4, e5 = fmt_metric(r3["ece"]), fmt_metric(r4["ece"]), fmt_metric(r5["ece"])
    if intervals_overlap(r3["ece_ci"], r4["ece_ci"]):
        cal_s = "did not change calibration detectably"
    elif r4["ece"] < r3["ece"]:
        cal_s = "improved calibration"
    else:
        cal_s = "worsened calibration"
    d_ts = cmp_formatted(r5["ece"], r3["ece"], fmt_metric)
    ts_s = (f"temperature scaling reduced the ECE of the served path to {e5}" if d_ts < 0 else
            f"temperature scaling did not reduce the ECE of the served path ({e5})")
    s1_ = (f"At a fixed checkpoint, PGS averaging {cal_s} (ECE {e4} versus {e3} for the served argmax path), "
           f"and {ts_s}.")
    verdict = {c["id"]: c["verdict"] for c in claims}["C4"]
    a5, a1 = fmt_metric(sel["S5"]["auroc"]), fmt_metric(sel["S1"]["auroc"])
    mi_defer_helps = _deferral_clause(results, sel["S5"])[0] > 0
    tail = {
        "SUPPORTED": ("so the PGS uncertainty signal is a usable basis for deferring uncertain reports to human "
                      "operators" if mi_defer_helps else
                      "so the PGS uncertainty signal ranks errors above chance, although deferring the 20% least "
                      "certain reports by it did not raise the accuracy of the remaining ones"),
        "QUALIFIED": "so a useful deferral signal is already available without PGS",
        "CONTRADICTED": "so the PGS uncertainty signal did not identify misrouted reports better than chance",
    }[verdict]
    s2_ = (f"As an error detector, the mutual information reached an AUROC of {a5}, compared with {a1} for "
           f"1 − max p of the served path, {tail}.")
    return f"{s1_} {s2_}"


def build_fill_values(results: Mapping[str, Any], claims: Sequence[Mapping[str, Any]],
                      deployed: str) -> tuple[dict[str, str], dict[str, str]]:
    """Return (fill_values, definitions) covering every key in REQUIRED_KEYS."""
    fv: dict[str, str] = {}
    defs: dict[str, str] = {}
    for r in results["calibration"]["rows"]:
        k = f"T5_{r['row']}"
        fv[f"{k}_ACC"] = fmt_metric(r["accuracy"])
        fv[f"{k}_CONF"] = fmt_metric(r["mean_confidence"])
        fv[f"{k}_ECE"] = fmt_metric_ci(r["ece"], *r["ece_ci"])
        fv[f"{k}_BRIER"] = fmt_metric(r["brier"])
        fv[f"{k}_NLL"] = fmt_metric(r["nll"])
        for m, d in (("ACC", "accuracy"), ("CONF", "mean top-label confidence"),
                     ("ECE", "top-label ECE (15 bins) with 95% bootstrap CI"),
                     ("BRIER", "multiclass Brier score (0–2)"), ("NLL", "negative log-likelihood (nats)")):
            defs[f"{k}_{m}"] = f"Table 5 {r['row']} ({r['description']}): {d}, test set"
    fv["T5_TEMP"] = fmt_temp(results["calibration"]["temperature"]["temperature"])
    defs["T5_TEMP"] = "temperature fitted on validation NLL of the served (pgs/argmax) probabilities"
    for r in results["paired"]["rows"]:
        k = f"T6_{r['row']}"
        fv[f"{k}_DACC"] = f"{fmt_pp(r['d_accuracy'])} {fmt_pp_ci(*r['d_accuracy_ci'])}"
        fv[f"{k}_DF1"] = f"{fmt_pp(r['d_macro_f1'])} {fmt_pp_ci(*r['d_macro_f1_ci'])}"
        fv[f"{k}_BC"] = f"{fmt_int(r['b'])} / {fmt_int(r['c'])}"
        fv[f"{k}_P"] = fmt_p(r["p_mcnemar"])
        fv[f"{k}_PHOLM"] = fmt_p(r["p_holm"])
        base = f"Table 6 {r['row']} (A = {r['label_a']}; B = {r['label_b']})"
        defs[f"{k}_DACC"] = f"{base}: accuracy A − B in pp with 95% paired bootstrap CI"
        defs[f"{k}_DF1"] = f"{base}: macro-F1 A − B in pp with 95% paired bootstrap CI"
        defs[f"{k}_BC"] = f"{base}: b = #(A correct, B wrong) / c = #(A wrong, B correct)"
        defs[f"{k}_P"] = f"{base}: McNemar exact two-sided p"
        defs[f"{k}_PHOLM"] = f"{base}: Holm-adjusted p within family '{r['family']}'"
    for r in results["selective"]["rows"]:
        k = f"T7_{r['row']}"
        fv[f"{k}_AUROC"] = fmt_metric_ci(r["auroc"], *r["auroc_ci"])
        fv[f"{k}_AURC"] = fmt_aurc(r["aurc"])
        fv[f"{k}_ACC80"] = fmt_metric(r["acc_at_80"])
        fv[f"{k}_COV90"] = fmt_cov(r["coverage_at_90"])
        base = f"Table 7 {r['row']} ({r['score']}: {r['description']}; correctness of {r['reference_predictions']})"
        defs[f"{k}_AUROC"] = f"{base}: AUROC for error detection with 95% bootstrap CI"
        defs[f"{k}_AURC"] = f"{base}: area under the risk–coverage curve ×100"
        defs[f"{k}_ACC80"] = f"{base}: accuracy of the 80% most certain samples"
        defs[f"{k}_COV90"] = f"{base}: largest coverage with accuracy ≥ 0.90 (0 if none)"
    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    pr = {r["row"]: r for r in results["paired"]["rows"]}
    sel = {r["score"]: r for r in results["selective"]["rows"]}
    fv["T1_ECE"] = f"{fmt_metric(cal['R3']['ece'])} vs. {fmt_metric(cal['R4']['ece'])}"
    defs["T1_ECE"] = ("ECE of the served argmax path (Table 5 R3) vs. PGS-averaged inference (Table 5 R4), "
                      "PGS checkpoint, test set — values in table-row order")
    # Label the p-value explicitly: Table 6 prints the raw and the Holm-adjusted p side by side, so a bare
    # "p =" in Table 1 would not match the raw column that a reader naturally compares it with.
    fv["T1_ENC"] = (f"Δmacro-F1 = {fmt_pp(pr['R1']['d_macro_f1'])} pp; "
                    f"{p_phrase(pr['R1']['p_holm'], 'Holm-adjusted p')}")
    defs["T1_ENC"] = (f"Table 6 R1 ({pr['R1']['label_a']} − {pr['R1']['label_b']}): macro-F1 difference and "
                      "Holm-adjusted McNemar p")
    fv["T1_AUROC"] = f"{fmt_metric(sel['S1']['auroc'])} vs. {fmt_metric(sel['S5']['auroc'])}"
    defs["T1_AUROC"] = ("Error-detection AUROC of 1 − max p of the served path (Table 7 R1, S1) vs. mutual "
                        "information (Table 7 R5, S5) — values in table-row order")
    fv["TXT_REPRO"] = txt_repro(results, deployed)
    fv["TXT_CALIBRATION"] = txt_calibration(results)
    fv["TXT_PAIRED"] = txt_paired(results)
    fv["TXT_SELECTIVE"] = txt_selective(results)
    fv["TXT_IMPLICATION"] = txt_implication(results, claims)
    fv["TXT_CONCLUSION"] = txt_conclusion(results, claims)
    for k, d in (("TXT_REPRO", "fragment continuing 'The exported records reproduce the accuracy and macro-F1 "
                               "values of Table 4 '"),
                 ("TXT_CALIBRATION", "calibration paragraph (Table 5, Figure 5)"),
                 ("TXT_PAIRED", "paired-test paragraph (Table 6)"),
                 ("TXT_SELECTIVE", "selective-prediction paragraph (Table 7, Figure 6)"),
                 ("TXT_IMPLICATION", "implication for the PGS claim and deferral recommendation"),
                 ("TXT_CONCLUSION", "conclusion sentences on calibration and error detection")):
        defs[k] = d
    missing = [k for k in REQUIRED_KEYS if k not in fv]
    assert not missing, f"fill_values lacks contract keys: {missing}"
    for k, v in fv.items():
        assert isinstance(v, str) and v and "\n" not in v, f"fill value {k} is empty or multi-line"
        assert "nan" not in v.lower().split(), f"fill value {k} contains nan: {v}"
    return {k: fv[k] for k in REQUIRED_KEYS}, {k: defs[k] for k in REQUIRED_KEYS}


# ---------------------------------------------------------------------------------------------
# Figures
# ---------------------------------------------------------------------------------------------


def _mpl():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    return matplotlib, plt


STYLE = {
    "font.family": "sans-serif",
    "font.sans-serif": ["Arial", "Helvetica", "Liberation Sans", "Nimbus Sans", "DejaVu Sans"],
    "font.size": 7.0,
    "axes.labelsize": 7.5,
    "axes.titlesize": 7.5,
    "xtick.labelsize": 6.5,
    "ytick.labelsize": 6.5,
    "legend.fontsize": 6.2,
    "axes.linewidth": 0.6,
    "xtick.major.width": 0.6,
    "ytick.major.width": 0.6,
    "xtick.major.size": 2.5,
    "ytick.major.size": 2.5,
    "xtick.minor.size": 1.5,
    "ytick.minor.size": 1.5,
    "xtick.minor.width": 0.4,
    "lines.linewidth": 1.1,
    "hatch.linewidth": 0.5,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "axes.unicode_minus": True,
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "savefig.facecolor": "white",
    "legend.frameon": False,
    "pdf.fonttype": 42,
    "mathtext.default": "regular",
}


def _check_png(path: Path, expected: tuple[int, int]) -> None:
    from PIL import Image

    with Image.open(path) as im:
        size = im.size
    if size != expected:
        raise AssertionError(f"{path.name}: {size[0]} x {size[1]} px, contract requires "
                             f"{expected[0]} x {expected[1]} px")


def _panel_label(ax: Any, text: str, x: float = -0.02) -> None:
    ax.text(x, 1.0, text, transform=ax.transAxes, fontsize=8, fontweight="bold", va="bottom", ha="right")


def _reliability_panel(ax: Any, ax_share: Any, row: Mapping[str, Any], title: str, first: bool) -> None:
    """Accuracy bars, hatched gap to mean confidence, diagonal, ECE box; sample share below."""
    b = row["bins"]
    lower = np.asarray(b["lower"], dtype=float)
    acc = np.asarray(b["accuracy"], dtype=float)
    conf = np.asarray(b["mean_confidence"], dtype=float)
    share = np.asarray(b["share"], dtype=float)
    ne = np.asarray(b["count"]) > 0
    w = 1.0 / N_BINS
    ax.bar(lower[ne], acc[ne], width=w, align="edge", color=OI["blue"], alpha=0.80,
           edgecolor="white", linewidth=0.4, zorder=2)
    gap_lo = np.minimum(acc, conf)
    gap_h = np.abs(acc - conf)
    ax.bar(lower[ne], gap_h[ne], bottom=gap_lo[ne], width=w, align="edge", facecolor=GAP_FACE,
           edgecolor=OI["vermillion"], linewidth=0.5, hatch="//////", zorder=3)
    ax.plot([0, 1], [0, 1], ls=(0, (3, 2)), color="0.2", lw=0.8, zorder=4)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.set_yticks(np.linspace(0, 1, 6))
    ax.tick_params(labelbottom=False)
    if first:
        ax.set_ylabel("Accuracy")
    ax.text(0.04, 0.96, f"{title}\nECE = {fmt_metric(row['ece'])}\n95% CI {fmt_ci(*row['ece_ci'])}",
            transform=ax.transAxes, va="top", ha="left", fontsize=6.3, linespacing=1.3,
            bbox=dict(boxstyle="round,pad=0.3", facecolor="white", edgecolor="0.75", linewidth=0.5), zorder=6)
    ax_share.bar(lower[ne], share[ne], width=w, align="edge", color="0.55", edgecolor="white", linewidth=0.4)
    ax_share.set_xlim(0, 1)
    top = float(np.nanmax(share)) if ne.any() else 1.0
    ax_share.set_ylim(0, top * 1.35)
    ax_share.set_yticks([0, top])
    ax_share.yaxis.set_major_formatter(lambda v, _pos: f"{v * 100:.0f}%")
    ax_share.tick_params(axis="y", labelsize=5.5, length=1.5, pad=1.5)
    ax_share.set_xticks(np.linspace(0, 1, 6))
    ax_share.set_xlabel("Confidence", labelpad=1.5)
    if first:
        ax_share.set_ylabel("Share", labelpad=2)


GAP_FACE = (0.835, 0.369, 0.0, 0.22)


def figure_reliability(results: Mapping[str, Any], path: Path) -> None:
    """Fig. 5: reliability diagrams (a) served argmax, (b) PGS-averaged, (c) argmax + TS."""
    matplotlib, plt = _mpl()
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    t = results["calibration"]["temperature"]["temperature"]
    panels = [
        ("(a)", cal["R3"], "Argmax (served path)"),
        ("(b)", cal["R4"], f"PGS-averaged (M = {_members(results)})"),
        ("(c)", cal["R5"], f"Argmax + temperature (T = {fmt_temp(t)})"),
    ]
    with matplotlib.rc_context(STYLE):
        fig = plt.figure(figsize=FIG_RELIABILITY_IN, dpi=DPI)
        gs = fig.add_gridspec(2, 3, height_ratios=[4.2, 1.0], left=0.062, right=0.992, bottom=0.135, top=0.885,
                              wspace=0.10, hspace=0.07)
        main_axes, share_axes = [], []
        for i, (lab, row, title) in enumerate(panels):
            ax = fig.add_subplot(gs[0, i], sharey=main_axes[0] if main_axes else None)
            ax_s = fig.add_subplot(gs[1, i], sharex=ax)
            _reliability_panel(ax, ax_s, row, title, first=(i == 0))
            if i:
                ax.tick_params(labelleft=False)
            _panel_label(ax, lab, x=-0.115 if i == 0 else -0.03)
            main_axes.append(ax)
            share_axes.append(ax_s)
        # common y-limits for the share strips so bars are comparable across panels
        top = max(a.get_ylim()[1] for a in share_axes)
        for i, a in enumerate(share_axes):
            a.set_ylim(0, top)
            if i:
                a.tick_params(labelleft=False)
        handles = [
            Patch(facecolor=OI["blue"], alpha=0.80, edgecolor="white", label="Accuracy per bin (15 bins)"),
            Patch(facecolor=GAP_FACE, edgecolor=OI["vermillion"], hatch="//////", linewidth=0.5,
                  label="Gap to mean confidence"),
            Line2D([0], [0], ls=(0, (3, 2)), color="0.2", lw=0.8, label="Perfect calibration"),
            Patch(facecolor="0.55", edgecolor="white", label="Share of samples per bin"),
        ]
        fig.legend(handles=handles, loc="upper center", ncol=4, bbox_to_anchor=(0.527, 1.0),
                   handlelength=1.6, columnspacing=2.0, borderaxespad=0.25)
        fig.savefig(path, dpi=DPI, facecolor="white")
        plt.close(fig)
    _check_png(path, FIG_RELIABILITY_PX)


def mi_floor(mi: np.ndarray) -> float:
    """Power of ten used as the lower edge of the log-scaled MI histogram.

    Exact zeros and values below the 0.5th percentile of the positive values (mostly
    floating-point cancellation noise in H(mean p) - mean H(p)) are drawn at this floor.
    """
    pos = mi[mi > 0]
    if pos.size == 0:
        return 1e-8
    return 10.0 ** math.floor(math.log10(float(np.quantile(pos, 0.005))))


def _ylim_clear_of_legend(fig: Any, ax: Any, legend: Any, lines: Sequence[tuple[np.ndarray, np.ndarray]],
                          lo: float, top: float, pad: float = 0.015) -> float:
    """Lower y-limit such that no plotted curve point runs under the (lower-left) legend.

    The legend is anchored in axes coordinates, so its height as a fraction of the axes does
    not depend on the y-limits; only the data have to move up. Returns ``lo`` unchanged when
    nothing collides, otherwise the largest multiple of 0.02 below the value that clears the
    legend top by ``pad`` (axes fraction). Assumes x-limits (0, 1).
    """
    bbox = legend.get_window_extent(fig.canvas.get_renderer()).transformed(ax.transAxes.inverted())
    x0, x1, frac = bbox.x0, bbox.x1, bbox.y1 + pad
    if not 0.0 < frac < 0.9:
        return lo
    under = [y[(x >= x0) & (x <= x1)] for x, y in lines]
    under = [u for u in under if u.size]
    if not under:
        return lo
    y_min = float(min(u.min() for u in under))
    if (y_min - lo) / (top - lo) > frac:
        return lo
    return math.floor((y_min - frac * top) / (1.0 - frac) * 50) / 50


def figure_selective(results: Mapping[str, Any], path: Path) -> None:
    """Fig. 6: (a) accuracy-coverage curves, (b) MI distribution for correct vs incorrect PGS predictions."""
    matplotlib, plt = _mpl()
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    sel = results["selective"]
    curves = sel["_curves"]
    rows = {r["score"]: r for r in sel["rows"]}
    cal = {r["row"]: r for r in results["calibration"]["rows"]}
    pgs_correct = ~sel["_incorrect"]["pgs"]
    mi = sel["_scores"]["S5"]
    with matplotlib.rc_context(STYLE):
        fig = plt.figure(figsize=FIG_SELECTIVE_IN, dpi=DPI)
        gs = fig.add_gridspec(1, 2, width_ratios=[1.12, 1.0], left=0.072, right=0.99, bottom=0.14,
                              top=0.955, wspace=0.24)
        ax = fig.add_subplot(gs[0, 0])
        n = pgs_correct.shape[0]
        k0 = max(1, int(math.ceil(0.01 * n)))
        spec = [
            ("S1", f"S1: 1 {MINUS} max p, served argmax", OI["blue"], "-", 1.3),
            ("S2", f"S2: 1 {MINUS} max p, PGS-averaged", OI["sky"], (0, (5, 1.5)), 1.1),
            ("S3", "S3: predictive entropy", OI["green"], (0, (1.2, 1.2)), 1.3),
            ("S5", "S5: mutual information", OI["orange"], (0, (5, 1.5, 1.2, 1.5)), 1.2),
        ]
        ymin_data = []
        for sid, label, color, ls, lw in spec:
            c = curves[sid]
            ax.plot(c.coverage[k0 - 1:], c.accuracy[k0 - 1:], color=color, ls=ls, lw=lw,
                    label=f"{label} (AUROC {fmt_metric(rows[sid]['auroc'])})", zorder=3)
            ymin_data.append(float(np.min(c.accuracy[k0 - 1:])))
        oc, oa = oracle_curve(pgs_correct)
        ax.plot(oc[k0 - 1:], oa[k0 - 1:], color="0.15", ls=(0, (3, 2)), lw=0.8,
                label="Oracle (PGS-averaged predictions)", zorder=2)
        full = cal["R4"]["accuracy"]
        ax.axhline(full, color="0.5", lw=0.7, ls=(0, (1, 1.5)), zorder=1,
                   label=f"Accuracy at full coverage, PGS-averaged ({fmt_metric(full)})")
        top = 1.004
        lo = math.floor((min(ymin_data + [full]) - 0.012) * 50) / 50
        ax.set_ylim(lo, top)
        ax.set_xlim(0, 1.0)
        ax.axvline(DEFER_COVERAGE, color="0.82", lw=0.6, zorder=0)
        ax.set_xlabel("Coverage")
        ax.set_ylabel("Accuracy of accepted reports")
        leg = ax.legend(loc="lower left", handlelength=2.8, borderaxespad=0.4, labelspacing=0.32, frameon=True,
                        framealpha=0.92, edgecolor="none", facecolor="white", borderpad=0.3)
        leg.set_zorder(5)
        plotted = [(curves[sid].coverage[k0 - 1:], curves[sid].accuracy[k0 - 1:]) for sid, *_ in spec]
        lo = _ylim_clear_of_legend(fig, ax, leg, plotted, lo, top)
        ax.set_ylim(lo, top)
        ax.text(DEFER_COVERAGE - 0.012, 1.0 - 0.012 * (top - lo) / 0.21, "80% coverage", rotation=90,
                fontsize=5.5, color="0.45", ha="right", va="top")
        _panel_label(ax, "(a)", x=-0.085)

        ax2 = fig.add_subplot(gs[0, 1])
        floor = mi_floor(mi)
        n_clipped = int(np.sum(mi < floor))
        clipped = np.maximum(mi, floor)
        hi = max(float(clipped.max()) * 1.05, floor * 10)
        bins = np.logspace(math.log10(floor), math.log10(hi), 40)
        peak = 0.0
        groups = [(pgs_correct, OI["blue"]), (~pgs_correct, OI["vermillion"])]
        for mask, color in groups:
            vals = clipped[mask]
            if vals.size == 0:
                continue
            weights = np.full(vals.shape, 1.0 / vals.size)
            h, _, _ = ax2.hist(vals, bins=bins, weights=weights, color=color, alpha=0.28, linewidth=0, zorder=2)
            ax2.hist(vals, bins=bins, weights=weights, histtype="step", color=color, linewidth=0.9, zorder=3)
            peak = max(peak, float(np.max(h)))
        ax2.set_xscale("log")
        ax2.set_xlim(bins[0], bins[-1])
        ax2.set_ylim(0, peak * 1.62)
        ymax_frac = peak * 1.08 / (peak * 1.62)
        for mask, color in groups:
            if mask.any():
                ax2.axvline(max(float(np.median(mi[mask])), floor), ymax=ymax_frac, color=color, lw=0.8,
                            ls=(0, (3, 1.5)), zorder=4)
        ax2.set_xlabel("Mutual information (nats, log scale)")
        ax2.set_ylabel("Fraction of predictions in group")
        ax2.yaxis.set_major_formatter(lambda v, _pos: fmt_fixed(v, 2))
        handles = [
            Patch(facecolor=matplotlib.colors.to_rgba(OI["blue"], 0.28), edgecolor=OI["blue"], linewidth=0.9,
                  label=f"Correct prediction (n = {fmt_int(int(pgs_correct.sum()))})"),
            Patch(facecolor=matplotlib.colors.to_rgba(OI["vermillion"], 0.28), edgecolor=OI["vermillion"],
                  linewidth=0.9, label=f"Incorrect prediction (n = {fmt_int(int((~pgs_correct).sum()))})"),
            Line2D([0], [0], color="0.3", lw=0.8, ls=(0, (3, 1.5)), label="Group median"),
        ]
        floor_s = _sci(floor)
        note = (f"MI < {floor_s} (incl. 0) drawn at {floor_s}: n = {fmt_int(n_clipped)}" if n_clipped
                else f"No MI value below {floor_s}")
        handles.append(Line2D([], [], lw=0, label=note))
        ax2.legend(handles=handles, loc="upper left", handlelength=1.8, borderaxespad=0.3, labelspacing=0.32)
        ax2.text(0.985, 0.975, f"AUROC (MI) = {fmt_metric(rows['S5']['auroc'])}\n95% CI {fmt_ci(*rows['S5']['auroc_ci'])}",
                 transform=ax2.transAxes, ha="right", va="top", fontsize=6.2, linespacing=1.3)
        _panel_label(ax2, "(b)", x=-0.10)
        fig.savefig(path, dpi=DPI, facecolor="white")
        plt.close(fig)
    _check_png(path, FIG_SELECTIVE_PX)


def _sci(x: float) -> str:
    """Power-of-ten label for matplotlib mathtext, e.g. "10$^{-9}$" or "2.5x10$^{-4}$" (Unicode minus)."""
    if not _finite(x) or x == 0:
        return "0"
    e = int(math.floor(math.log10(abs(x))))
    m = x / 10 ** e
    ms = f"{m:.1f}"
    if ms == "10.0":
        ms, e = "1.0", e + 1
    exp = str(e).replace("-", MINUS)
    return f"10$^{{{exp}}}$" if ms == "1.0" else f"{ms}×10$^{{{exp}}}$"


# ---------------------------------------------------------------------------------------------
# Tables and reports
# ---------------------------------------------------------------------------------------------


def _md_table(header: Sequence[str], rows: Iterable[Sequence[str]]) -> str:
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join("---" for _ in header) + "|"]
    lines += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    return "\n".join(lines)


def write_tables(results: Mapping[str, Any], fv: Mapping[str, str], defs: Mapping[str, str],
                 out: Path, deployed: str) -> None:
    tdir = out / "tables"
    tdir.mkdir(parents=True, exist_ok=True)
    cal = results["calibration"]
    pr = results["paired"]
    sel = results["selective"]
    rep = results["reproduction"]

    pd.DataFrame([{
        "row": r["row"], "description": r["description"], "checkpoint": r["checkpoint"], "mode": r["mode"],
        "temperature_scaled": r["temperature_scaled"], "n": r["n"], "accuracy": r["accuracy"],
        "mean_confidence": r["mean_confidence"], "ece": r["ece"], "ece_ci_lo": r["ece_ci"][0],
        "ece_ci_hi": r["ece_ci"][1], "brier": r["brier"], "brier_ci_lo": r["brier_ci"][0],
        "brier_ci_hi": r["brier_ci"][1], "nll": r["nll"], "nll_ci_lo": r["nll_ci"][0], "nll_ci_hi": r["nll_ci"][1],
    } for r in cal["rows"]]).to_csv(tdir / "table5_calibration.csv", index=False)
    bins_rows = []
    for r in cal["rows"]:
        b = r["bins"]
        for i in range(N_BINS):
            bins_rows.append({"row": r["row"], "bin": i, "lower": b["lower"][i], "upper": b["upper"][i],
                              "count": b["count"][i], "share": b["share"][i], "accuracy": b["accuracy"][i],
                              "mean_confidence": b["mean_confidence"][i]})
    pd.DataFrame(bins_rows).to_csv(tdir / "reliability_bins.csv", index=False)
    pd.DataFrame([{
        "row": r["row"], "family": r["family"], "system_a": r["label_a"], "system_b": r["label_b"], "n": r["n"],
        "accuracy_a": r["accuracy_a"], "accuracy_b": r["accuracy_b"], "d_accuracy_pp": 100 * r["d_accuracy"],
        "d_accuracy_ci_lo_pp": 100 * r["d_accuracy_ci"][0], "d_accuracy_ci_hi_pp": 100 * r["d_accuracy_ci"][1],
        "macro_f1_a": r["macro_f1_a"], "macro_f1_b": r["macro_f1_b"], "d_macro_f1_pp": 100 * r["d_macro_f1"],
        "d_macro_f1_ci_lo_pp": 100 * r["d_macro_f1_ci"][0], "d_macro_f1_ci_hi_pp": 100 * r["d_macro_f1_ci"][1],
        "b": r["b"], "c": r["c"], "p_mcnemar": r["p_mcnemar"], "p_holm": r.get("p_holm"),
        "top1_disagreements": r["top1_disagreements"],
    } for r in pr["rows"] + pr["supplementary"]]).to_csv(tdir / "table6_paired.csv", index=False)
    pd.DataFrame([{
        "row": r["row"], "score": r["score"], "description": r["description"],
        "reference_predictions": r["reference_predictions"], "auroc": r["auroc"], "auroc_ci_lo": r["auroc_ci"][0],
        "auroc_ci_hi": r["auroc_ci"][1], "aurc_x100": r["aurc_x100"], "acc_at_80": r["acc_at_80"],
        "k_at_80": r["k_at_80"], "coverage_at_90": r["coverage_at_90"], "k_at_90": r["k_at_90"],
        "full_coverage_accuracy": r["full_coverage_accuracy"], "n_errors": r["n_errors"],
    } for r in sel["rows"]]).to_csv(tdir / "table7_selective.csv", index=False)
    pd.DataFrame(sel["curves_on_grid"]).to_csv(tdir / "selective_curves.csv", index=False)
    pd.DataFrame(sel["paired_auroc_differences"]).assign(
        d_auroc_ci_lo=lambda d: d["d_auroc_ci"].map(lambda x: x[0]),
        d_auroc_ci_hi=lambda d: d["d_auroc_ci"].map(lambda x: x[1])).drop(columns="d_auroc_ci").to_csv(
        tdir / "supplementary_auroc_differences.csv", index=False)
    pd.DataFrame([{
        "file": r["file"], "pair": r["pair"], "checkpoint": r["checkpoint"], "mode": r["mode"],
        "accuracy": r["accuracy"], "macro_f1": r["macro_f1"],
        "expected_accuracy": " or ".join(f"{x:.4f}" for x in r["checks"].get("accuracy", {}).get("expected", [])),
        "expected_macro_f1": " or ".join(f"{x:.4f}" for x in r["checks"].get("macro_f1", {}).get("expected", [])),
        "reproduces_computed": r["reproduces_computed"], "reproduces_manifest": r["reproduces_manifest"],
        "tree_count": r["tree_count"],
    } for r in rep["rows"]]).to_csv(tdir / "table4_reproduction.csv", index=False)
    pd.DataFrame([{"key": k, "value": fv[k], "definition": defs[k]} for k in ("T1_ECE", "T1_ENC", "T1_AUROC")]
                 ).to_csv(tdir / "table1_additions.csv", index=False)

    t = cal["temperature"]
    md = [f"# Tables — per-sample uncertainty analysis ({deployed}, test n = {fmt_int(cal['rows'][0]['n'])})", "",
          "All values are exactly the strings written to `fill_values.json`. Bootstrap: "
          f"{fmt_int(results['meta']['bootstrap_resamples'])} resamples, seed {results['meta']['seed']}; "
          "95% percentile intervals.", ""]
    md += ["## Table 4 check — export vs. manuscript (test set)", "",
           _md_table(["Pair", "Ckpt / mode", "Accuracy", "Expected", "Macro-F1", "Expected", "Reproduces"],
                     [[r["pair"], f"{r['checkpoint']} / {r['mode']}", fmt_metric(r["accuracy"]),
                       " or ".join(fmt_metric(x) for x in r["checks"].get("accuracy", {}).get("expected", [])) or "—",
                       fmt_metric(r["macro_f1"]),
                       " or ".join(fmt_metric(x) for x in r["checks"].get("macro_f1", {}).get("expected", [])) or "—",
                       {True: "yes", False: "NO", None: "—"}[r["reproduces_computed"]]] for r in rep["rows"]]),
           "", f"Top-1 disagreements pgs/argmax vs. pgs/pgs: {fmt_int(rep['top1_disagreements_pgs_vs_argmax'])} "
               "(manuscript: 130).", ""]
    md += ["## Table 5 — calibration (deployed pair, test set)", "",
           _md_table(["Row", "Configuration", "Accuracy", "Mean conf.", "ECE [95% CI]", "Brier", "NLL"],
                     [[r["row"], r["description"], fv[f"T5_{r['row']}_ACC"], fv[f"T5_{r['row']}_CONF"],
                       fv[f"T5_{r['row']}_ECE"], fv[f"T5_{r['row']}_BRIER"], fv[f"T5_{r['row']}_NLL"]]
                      for r in cal["rows"]]),
           "", f"Temperature T = {fv['T5_TEMP']} (validation NLL {fmt_metric(t['val_nll_before'])} → "
               f"{fmt_metric(t['val_nll_after'])}; validation ECE {fmt_metric(t['val_ece_before'])} → "
               f"{fmt_metric(t['val_ece_after'])}).", ""]
    md += ["## Table 6 — paired tests (test set)", "",
           _md_table(["Row", "A", "B", "ΔAccuracy, pp [95% CI]", "ΔMacro-F1, pp [95% CI]", "b / c", "McNemar p",
                      "Holm p"],
                     [[r["row"], r["label_a"], r["label_b"], fv[f"T6_{r['row']}_DACC"], fv[f"T6_{r['row']}_DF1"],
                       fv[f"T6_{r['row']}_BC"], fv[f"T6_{r['row']}_P"], fv[f"T6_{r['row']}_PHOLM"]]
                      for r in pr["rows"]]),
           "", "Families for Holm: R1–R3 (encoders), R4–R5 (inference modes). b = #(A correct, B wrong), "
               "c = #(A wrong, B correct).", ""]
    s = pr["supplementary"][0]
    md += [f"Supplementary (not in Table 6): {s['label_a']} − {s['label_b']}: Δaccuracy = {fmt_pp(s['d_accuracy'])} "
           f"{fmt_pp_ci(*s['d_accuracy_ci'])} pp, Δmacro-F1 = {fmt_pp(s['d_macro_f1'])} {fmt_pp_ci(*s['d_macro_f1_ci'])} "
           f"pp, b / c = {fmt_int(s['b'])} / {fmt_int(s['c'])}, McNemar {p_phrase(s['p_mcnemar'])}.", ""]
    md += ["## Table 7 — selective prediction (deployed pair, PGS checkpoint, test set)", "",
           _md_table(["Row", "Score", "AUROC [95% CI]", "AURC ×100", "Acc. @ 80% cov.", "Cov. @ 90% acc."],
                     [[r["row"], f"{r['score']}: {r['description']}", fv[f"T7_{r['row']}_AUROC"],
                       fv[f"T7_{r['row']}_AURC"], fv[f"T7_{r['row']}_ACC80"], fv[f"T7_{r['row']}_COV90"]]
                      for r in sel["rows"]]),
           "", "S1 is scored against the pgs/argmax predictions; S2–S6 against the pgs/pgs predictions.", "",
           "Paired bootstrap ΔAUROC (supplementary): " + "; ".join(
               f"{d['a']} − {d['b']} = {fmt_fixed(d['d_auroc'], 4, plus=True)} {fmt_ci(*d['d_auroc_ci'])}"
               for d in sel["paired_auroc_differences"]) + ".", ""]
    md += ["## Table 1 additions", "",
           _md_table(["Key", "Value", "Definition"], [[k, fv[k], defs[k]] for k in ("T1_ECE", "T1_ENC", "T1_AUROC")]),
           ""]
    md += ["## Generated prose", ""] + [f"**{k}.** {fv[k]}\n" for k in REQUIRED_KEYS if k.startswith("TXT_")]
    (out / "tables.md").write_text("\n".join(md), encoding="utf-8")


def write_claims(claims: Sequence[Mapping[str, Any]], results: Mapping[str, Any], out: Path,
                 manuscript: Path | None) -> None:
    meta = results["meta"]
    lines = ["# Manuscript claims check (CONTRACT.md §5)", "",
             f"Predictions: `{meta['preds_dir']}`; deployed pair `{meta['deployed']}`; test n = "
             f"{fmt_int(meta['n_test'])}; bootstrap {fmt_int(meta['bootstrap_resamples'])} resamples (seed "
             f"{meta['seed']}); generated {meta['created_utc']}.",
             "Quoted sentences are " + (f"extracted from `{manuscript.name}`." if manuscript else
                                        "the wording of the submitted revision (run with `--manuscript` to "
                                        "re-extract them from the current DOCX)."), "",
             _md_table(["Claim", "Verdict", "Criterion"],
                       [[f"{c['id']} {c['title']}", f"**{c['verdict']}**", c["criterion"]] for c in claims]), ""]
    for c in claims:
        src = CLAIM_SOURCES[c["id"]]
        quotes = src["quotes"]
        if manuscript is not None:
            try:
                found = manuscript_sentences(manuscript, src["phrases"])
                quotes = found or quotes
            except Exception as exc:  # pragma: no cover - depends on the DOCX at hand
                LOG.warning("could not read %s: %s", manuscript, exc)
        lines += [f"## {c['id']} — {c['title']}", "",
                  f"**Verdict: {c['verdict']}**", "",
                  f"*Criterion:* {c['criterion']}", "",
                  *([f"*Evidence:* {c['evidence']}", ""] if "\n" not in c["evidence"]
                    else ["*Evidence:*", "", c["evidence"], ""]),
                  f"*Manuscript sentence(s) ({c['sections']}):*", ""]
        for q in quotes:
            lines += [f"> {q}", ""]
        if c["verdict"] == "SUPPORTED":
            lines += ["*Action:* no revision required."]
            if c.get("refinement"):
                lines += ["", f"*Recommended refinement:* {c['refinement']}"]
        else:
            if not c.get("suggestion"):   # never print a bare "None"
                lines += ["*Action:* check the evidence above before relying on the corresponding text."]
            elif c.get("suggestion_is_instruction"):
                lines += [f"*Action:* {c['suggestion']}"]
            else:
                lines += [f"*Action:* revise the sentence(s) above. Suggested wording: \"{c['suggestion']}\""]
        lines += [""]
    (out / "claims_check.md").write_text("\n".join(lines), encoding="utf-8")


def _jsonable(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {str(k): _jsonable(v) for k, v in obj.items() if not str(k).startswith("_")}
    if isinstance(obj, (list, tuple)):
        return [_jsonable(v) for v in obj]
    if isinstance(obj, np.ndarray):
        return _jsonable(obj.tolist())
    if isinstance(obj, (np.bool_, bool)):
        return bool(obj)
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (float, np.floating)):
        return float(obj) if math.isfinite(float(obj)) else None
    if isinstance(obj, Path):
        return obj.as_posix()
    return obj


# ---------------------------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------------------------


def run(preds_dir: Path, out_dir: Path, deployed: str = DEFAULT_DEPLOYED,
        encoders: Sequence[str] = DEFAULT_ENCODERS, n_boot: int = 2000, seed: int = 42,
        manuscript: Path | None = None, make_figures: bool = True) -> dict[str, Any]:
    """Run the whole stage-2 analysis and write every output file. Returns the results dict."""
    t0 = time.perf_counter()
    if "__" not in deployed:
        raise SchemaError(f"--deployed must look like image__text, got {deployed!r}")
    if len(encoders) != 3:
        raise SchemaError("--encoders must name exactly three image encoders (R1-R3)")
    manifest, files = load_prediction_dir(preds_dir)
    img, txt = deployed.split("__", 1)
    encoder_pairs = [f"{e}__{txt}" for e in encoders]
    if encoder_pairs[0] != deployed:
        LOG.warning("first encoder %s is not the deployed image encoder %s", encoders[0], img)
    needed = [(deployed, c, m) for c in ("cb", "pgs") for m in ("argmax", "pgs")]
    needed += [(p, "pgs", "pgs") for p in encoder_pairs]
    missing = [k for k in needed if k not in files]
    if missing:
        raise SchemaError("required prediction files missing: " + ", ".join(
            f"preds__{p}__{c}__{m}.csv.gz" for p, c, m in missing))
    ref_df = files[(deployed, "pgs", "pgs")].frame
    ref_test = ref_df.loc[ref_df["split"] == "test", ["split", "row_id", "label"]].reset_index(drop=True)
    test_views = {k: make_view(pf, "test", ref_test) for k, pf in files.items()}
    val_served = make_view(files[(deployed, "pgs", "argmax")], "val")
    n_test = ref_test.shape[0]
    t_load = time.perf_counter()
    LOG.info("loaded %d files (n_test = %d) in %.1f s", len(files), n_test, t_load - t0)

    boot = Bootstrap(n_test, n_boot, seed)
    t_boot = time.perf_counter()
    dep_views = {(c, m): test_views[(deployed, c, m)] for c in ("cb", "pgs") for m in ("argmax", "pgs")}
    results: dict[str, Any] = {
        "meta": {
            "created_utc": _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "script": "analyze_uncertainty.py", "preds_dir": preds_dir.as_posix(), "deployed": deployed,
            "encoder_pairs": encoder_pairs, "n_test": n_test, "n_val": val_served.n,
            "bootstrap_resamples": n_boot, "seed": seed,
            "bootstrap_scheme": ("numpy.random.default_rng(seed).integers(0, N, size=(B, N)) drawn once and "
                                 "shared by all metrics; paired comparisons reuse the same resamples; 95% "
                                 "percentile intervals (numpy.percentile, linear interpolation)"),
            "ece": "top-label, 15 equal-width bins (i/15, (i+1)/15], confidence = max p",
            "temperature_scaling": "T minimises validation NLL of softmax(log clip(p, 1e-12, 1) / T), bounded "
                                   "search T in [0.05, 20]; fitted on pgs/argmax validation probabilities",
            "selective": "ascending score (stable, ties by row order); AURC = mean_k errors(k)/k; accuracy at "
                         "coverage round(0.8 N); coverage at 90% accuracy = max k/N with accuracy(k) >= 0.90",
            "versions": {"python": platform.python_version(), "numpy": np.__version__, "pandas": pd.__version__,
                         "scipy": __import__("scipy").__version__,
                         "sklearn": __import__("sklearn").__version__},
            "file_warnings": {pf.path.name: pf.warnings for pf in files.values() if pf.warnings},
        },
    }
    results["reproduction"] = analyse_reproduction(manifest, files, test_views, deployed)
    results["calibration"] = analyse_calibration(dep_views, val_served, boot)
    results["calibration"]["all_files_point_estimates"] = [
        {"file": k[0] + "__" + k[1] + "__" + k[2], "accuracy": float(v.correct.mean()),
         "mean_confidence": float(v.confidence.mean()), "ece": ece_score(v.probs, v.label, v.pred),
         "brier": float(brier_per_sample(v.probs, v.label).mean()),
         "nll": float(nll_per_sample(v.probs, v.label).mean())} for k, v in sorted(test_views.items())]
    t_cal = time.perf_counter()
    results["paired"] = analyse_paired(test_views, deployed, encoder_pairs, boot)
    t_pair = time.perf_counter()
    results["selective"] = analyse_selective(dep_views[("pgs", "argmax")], dep_views[("pgs", "pgs")], boot)
    t_sel = time.perf_counter()
    claims = evaluate_claims(results)
    results["claims"] = claims
    fv, defs = build_fill_values(results, claims, deployed)
    results["fill_value_definitions"] = defs

    out_dir.mkdir(parents=True, exist_ok=True)
    if make_figures:
        figure_reliability(results, out_dir / "fig_reliability.png")
        figure_selective(results, out_dir / "fig_selective.png")
    t_fig = time.perf_counter()
    write_tables(results, fv, defs, out_dir, deployed)
    write_claims(claims, results, out_dir, manuscript)
    results["meta"]["timing_seconds"] = {
        "load": t_load - t0, "bootstrap_setup": t_boot - t_load, "calibration": t_cal - t_boot,
        "paired": t_pair - t_cal, "selective": t_sel - t_pair, "figures": t_fig - t_sel,
        "total": time.perf_counter() - t0,
    }
    (out_dir / "fill_values.json").write_text(json.dumps(fv, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out_dir / "results.json").write_text(json.dumps(_jsonable(results), ensure_ascii=False, indent=2) + "\n",
                                          encoding="utf-8")
    LOG.info("wrote results to %s in %.1f s", out_dir, time.perf_counter() - t0)
    return results


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0],
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--preds", type=Path, required=True, help="per_sample_predictions directory (stage 1)")
    p.add_argument("--out", type=Path, required=True, help="results directory to write")
    p.add_argument("--deployed", default=DEFAULT_DEPLOYED, help="deployed pair image__text (default %(default)s)")
    p.add_argument("--encoders", default=",".join(DEFAULT_ENCODERS),
                   help="three image encoders for Table 6 R1-R3, first = deployed (default %(default)s)")
    p.add_argument("--bootstrap", type=int, default=2000, help="bootstrap resamples (default %(default)s)")
    p.add_argument("--seed", type=int, default=42, help="bootstrap seed (default %(default)s)")
    p.add_argument("--manuscript", type=Path, default=None,
                   help="optional DOCX, read-only, to quote the current claim sentences in claims_check.md")
    p.add_argument("--no-figures", action="store_true", help="skip the two PNG figures")
    p.add_argument("-v", "--verbose", action="store_true")
    return p


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(levelname)s %(name)s: %(message)s")
    encoders = [e.strip() for e in args.encoders.split(",") if e.strip()]
    try:
        results = run(args.preds, args.out, args.deployed, encoders, args.bootstrap, args.seed,
                      args.manuscript, make_figures=not args.no_figures)
    except SchemaError as exc:
        LOG.error("%s", exc)
        return 2
    verdicts = ", ".join(f"{c['id']}={c['verdict']}" for c in results["claims"])
    LOG.info("claims: %s", verdicts)
    return 0


if __name__ == "__main__":
    sys.exit(main())
